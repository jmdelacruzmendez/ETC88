#!/usr/bin/env python3
"""Genera el DOC DE TRABAJO en .docx para la reunión de coordinadores ETC 88.
Material vivo: tablas, enumeraciones, negritas, cursivas + logo del ETC.
Pensado para escribir arriba en vivo (cada sección tiene espacio para notas).
Salida: entrega_diseno/DOC_TRABAJO_REUNION_COORDS.docx
"""
import os, re
import cairosvg
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(REPO, 'entrega_diseno')
os.makedirs(OUT_DIR, exist_ok=True)
LOGO_SVG = os.path.join(REPO, 'design', 'logo_eteciano.svg')
LOGO_PNG = os.path.join(REPO, 'design', 'logo_eteciano.png')

# Convertir SVG → PNG para que python-docx pueda embeberlo
cairosvg.svg2png(url=LOGO_SVG, write_to=LOGO_PNG, output_width=900)

# ===== Estilos (Dirección A) =====
AZUL  = RGBColor(0x1B, 0x3A, 0x5C)
ARENA = RGBColor(0xC0, 0xA0, 0x6C)
CORAL = RGBColor(0xE3, 0x6C, 0x4F)
TINTA = RGBColor(0x1A, 0x1A, 0x1A)
GRIS  = RGBColor(0x66, 0x66, 0x66)

FUENTE = 'Calibri'

# ===== Helpers =====
def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_runs(para, text, size=11, color=TINTA):
    """Procesa **bold**, *italic* y _underline_ en un texto plano."""
    if not text:
        para.add_run('')
        return
    pattern = r'(\*\*[^*]+?\*\*|\*[^*\n]+?\*|_[^_\n]+?_)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = para.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            r = para.add_run(part[1:-1]); r.italic = True
        elif part.startswith('_') and part.endswith('_'):
            r = para.add_run(part[1:-1]); r.underline = True
        else:
            r = para.add_run(part)
        r.font.name = FUENTE
        r.font.size = Pt(size)
        r.font.color.rgb = color

def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True
    r.font.name = FUENTE; r.font.size = Pt(20); r.font.color.rgb = AZUL
    # Línea bajo H1
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), 'C0A06C')
    pBdr.append(bottom)
    pPr.append(pBdr)

def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True
    r.font.name = FUENTE; r.font.size = Pt(14); r.font.color.rgb = AZUL

def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True
    r.font.name = FUENTE; r.font.size = Pt(11); r.font.color.rgb = AZUL

def para(doc, text=None, italic=False, color=TINTA, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    if text is None:
        return p
    add_runs(p, text, size=size, color=color)
    if italic and p.runs:
        for r in p.runs: r.italic = True
    return p

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Cm(0.75 + 0.6*level)
    p.paragraph_format.space_after = Pt(1)
    add_runs(p, text, size=11)
    return p

def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(2)
    add_runs(p, text, size=11)
    return p

def callout(doc, text):
    """Cuadro de cita / nota con borde y fondo arena suave."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    set_cell_bg(cell, 'F7F1E3')
    cell.width = Cm(16)
    # borde izquierdo coral
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, color, sz in [('left','E36C4F','24'), ('top','E8D4A8','4'),
                             ('right','E8D4A8','4'), ('bottom','E8D4A8','4')]:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), sz)
        b.set(qn('w:color'), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)
    cell.text = ''
    p = cell.paragraphs[0]
    add_runs(p, text, size=11, color=AZUL)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def linea_para_escribir(doc, label=''):
    """Línea con underscore para escribir a mano arriba o llenar en Word."""
    p = doc.add_paragraph()
    if label:
        r = p.add_run(label + '  ')
        r.font.name = FUENTE; r.font.size = Pt(11); r.italic = True; r.font.color.rgb = GRIS
    r = p.add_run('_' * 60)
    r.font.name = FUENTE; r.font.size = Pt(11); r.font.color.rgb = GRIS

def espacio_notas(doc, titulo='Notas / decisiones en vivo', lineas=4):
    h3(doc, titulo)
    for _ in range(lineas):
        linea_para_escribir(doc)

def make_table(doc, headers, rows, col_widths_cm=None, header_bg='1B3A5C', header_fg='FFFFFF'):
    t = doc.add_table(rows=len(rows)+1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Headers
    for j, h in enumerate(headers):
        c = t.cell(0, j)
        set_cell_bg(c, header_bg)
        c.text = ''
        p = c.paragraphs[0]
        r = p.add_run(h); r.bold = True
        r.font.name = FUENTE; r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(header_fg)
    # Rows
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ''
            p = c.paragraphs[0]
            add_runs(p, str(val), size=10)
    # Column widths
    if col_widths_cm:
        for j, w in enumerate(col_widths_cm):
            for r_ in t.rows:
                r_.cells[j].width = Cm(w)
    return t

# ===== Documento =====
doc = Document()

# Margins
for s in doc.sections:
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

# ---------- PORTADA ----------
# Logo centrado
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(LOGO_PNG, width=Inches(2.2))
p.paragraph_format.space_after = Pt(12)

# Título
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ETC 88'); r.bold = True
r.font.name = FUENTE; r.font.size = Pt(32); r.font.color.rgb = AZUL

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Documento de Trabajo'); r.bold = True
r.font.name = FUENTE; r.font.size = Pt(22); r.font.color.rgb = AZUL

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Reunión de Coordinadores · 4 de junio de 2026')
r.font.name = FUENTE; r.font.size = Pt(14); r.italic = True; r.font.color.rgb = ARENA

doc.add_paragraph().paragraph_format.space_after = Pt(20)

# Lema + hilo
callout(doc,
    '**Lema eteciano:** *"Siempre amigos"* · Jn 15,15\n'
    '**Lema del retiro:** *"Donde está tu tesoro, allí estará tu corazón"* · Mt 6,21\n'
    '**Hilo espiritual:** El corazón y el tesoro — Mt 6,21\n'
    '**Retiro:** 4–6 de septiembre de 2026 · Casa de Retiro «La Ceiba del Salado», Higüey')

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Material vivo · pensado para escribir arriba y trabajar en vivo')
r.font.name = FUENTE; r.font.size = Pt(11); r.italic = True; r.font.color.rgb = GRIS

# Convocados
doc.add_paragraph().paragraph_format.space_after = Pt(6)
h2(doc, 'Convocados a la reunión')
bullet(doc, '**Co-Directores:** Juan Manuel de la Cruz · Jean Carlo de la Cruz')
bullet(doc, '**Asesores del Retiro:** Laura · Tomás')
bullet(doc, '**Coordinadores de Guías:** Priscilla + Camila')
bullet(doc, '**Coordinadores de Cocina:** Paloma + Jhonnito')
bullet(doc, '**Coordinador de Música:** José Ángel Tusen')

doc.add_page_break()

# ---------- §1 AGENDA ----------
h1(doc, '1 · Agenda de la reunión')
items = [
    'Apertura + oración (Co-Dir)',
    'Equipo cerrado — quién es quién',
    'Cómo gestionas tu primera reunión con tu equipo',
    'Roles y Responsabilidades del Coordinador (5 principios)',
    'El tablero — qué hay y cómo se usa por área',
    'Qué requerimos de cada equipo + Roadmap por área',
    'Materiales del retiro — por área',
    'Pagos mensuales del equipo — cómo se cobran',
    'Disciplina y asistencia',
    'Cierre + compromisos + oración',
]
for it in items:
    numbered(doc, it)

para(doc)
espacio_notas(doc, 'Ajustes a la agenda en vivo', 3)

doc.add_page_break()

# ---------- §2 EQUIPO ----------
h1(doc, '2 · Equipo del ETC 88')
para(doc, 'Composición del equipo con sus conteos y coordinaciones por área:')
make_table(doc,
    headers=['Área', 'Personas', 'Coordinación / Notas'],
    rows=[
        ['Directores', '2', 'Juan Manuel + Jean Carlo'],
        ['Asesores del Retiro', '3', 'Laura · Tomás · Frank'],
        ['Asesores Espirituales transversales', '2', 'Padre Paul Ramírez + Sor Angelina Lebrón'],
        ['Guías', '14', '7 parejas color-coded · Coords: Priscilla + Camila'],
        ['Cocina', '21', 'Coords: Paloma + Jhonnito · Asesoras: Petra + Johanny'],
        ['Música', '6', 'Coord: José Ángel Tusen'],
        ['Asesores de comunidad', '3', 'La Vega (Leticia) · SD (Marleny + Sandrita)'],
    ],
    col_widths_cm=[5, 2.5, 9])

para(doc)
h3(doc, '5 equipos auxiliares en construcción')
para(doc, 'Responsables por definir. Salen de esta reunión con un primer nombre por equipo si es posible.')
make_table(doc,
    headers=['Equipo auxiliar', 'Función', 'Responsable [POR DEFINIR]'],
    rows=[
        ['Recaudación / Donaciones', 'Coordina cartas, contactos y campañas para cubrir el costo del retiro', '_________________'],
        ['Guagua / Transporte', 'Cotiza, contrata y gestiona transporte SPM → Higüey', '_________________'],
        ['Profondo', 'Actividad pro-fondos (31-jul → 2-ago)', '_________________'],
        ['Intersección (diáspora)', 'Enlace con etecianos en el exterior', '_________________'],
        ['Finanzas / Tesorería', 'Control de pagos del equipo, presupuesto, reporte', '_________________'],
    ],
    col_widths_cm=[5, 7, 4.5])

espacio_notas(doc, 'Notas del equipo en vivo', 4)
doc.add_page_break()

# ---------- §3 CÓMO GESTIONAS TU 1RA REUNIÓN ----------
h1(doc, '3 · Cómo gestionas tu primera reunión con tu equipo')
para(doc, 'Cada coord tendrá pronto su primera reunión con su equipo de área. Para que arranque bien, estos 7 pasos:')
for paso in [
    '**Presenta a los Asesores Espirituales transversales** (Padre Paul + Sor Angelina) — están para acompañar al equipo durante todo el camino, no solo el día del retiro.',
    '**Transmite los 5 principios** (§4) **sin reinterpretarlos.** El coord es el primer canal del director hacia el equipo.',
    '**Presentación de cada miembro** — que se conozcan; cada área llega con gente nueva y con gente que ya ha servido antes.',
    '**Repaso del calendario** del ETC 88 (F1 a F5, Profondo, Convivencia, Ensayo General, Retiro).',
    '**Asigna lectura** del material que tu equipo necesita: Guía de Guías 88, Anexo de Cocina 88, Anexo de Música 88.',
    '**Fija la próxima reunión** (lugar + fecha + agenda) antes de cerrar.',
    '**Cierra con oración** — el ETC es ante todo una obra espiritual.',
]:
    numbered(doc, paso)
para(doc)
callout(doc, '**Tono:** descriptivo, no normativo. El coord acompaña, no manda. La guía es una guía, no un reglamento.')

espacio_notas(doc, 'Cosas para añadir a estos 7 pasos', 3)
doc.add_page_break()

# ---------- §4 ROLES Y RESPONSABILIDADES ----------
h1(doc, '4 · Roles y Responsabilidades del Coordinador')

h2(doc, 'El rol del coordinador')
para(doc, 'El coord **acompaña a su equipo de área**: es la voz del director hacia su gente, y la voz de su gente hacia el director. Da ejemplo (asiste a reuniones, llega a tiempo, prepara su material), transmite sin reinterpretar, escala dudas al Co-Dir, y cuida que cada miembro se sienta visto.')

h2(doc, 'Los 5 principios de comunicación del ETC 88')
for i, (titulo, desc) in enumerate([
    ('Enfocados en el participante', 'Todo lo que hacemos es para él. Cuando hay dudas, la pregunta es: *¿qué es mejor para el participante?*'),
    ('Transmitir, no reinterpretar', 'Los mensajes del director bajan **tal cual** al equipo. Si hay duda, se pregunta; no se reinterpreta.'),
    ('Un solo equipo, funciones distintas', 'Guías, cocina, música, asesores y directores somos un equipo. Cuando un área pide ayuda, el resto responde.'),
    ('Servicio y sacrificio antes que protagonismo', '**Jn 13** (Lavatorio de los pies). Vale para todas las áreas, no solo cocina.'),
    ('Coordinarse y ensayar de antemano', 'La fluidez se prepara, no se improvisa.'),
], start=1):
    h3(doc, f'{i}. {titulo}')
    para(doc, desc)

h2(doc, 'Lo que el coord NO hace')
bullet(doc, 'No improvisa decisiones que son del director.')
bullet(doc, 'No filtra preguntas que su gente tiene para el director.')
bullet(doc, 'No carga el peso solo: pide apoyo a los asesores y al Co-Dir.')

espacio_notas(doc, 'Decisiones y ajustes en vivo', 4)
doc.add_page_break()

# ---------- §5 EL TABLERO ----------
h1(doc, '5 · El tablero — qué hay y cómo se usa por área')
bullet(doc, 'El tablero + las hojas de Equipo y Control de Pagos = **fuente de verdad** del estado del equipo.')
bullet(doc, '**Cocina** mira las **alergias y condiciones** antes de cerrar el menú.')
bullet(doc, '**Guías** cuidan los **lazos / parejas / hermanos** — no asignar un participante al PG de su hermano o pareja del equipo.')
bullet(doc, '**Música** consulta el repertorio de referencia y el calendario de ensayos.')
bullet(doc, '**Cada dato tiene dueño** — si tu gente cambia datos, los cambia en el formulario; el sistema regenera, nadie edita el tablero a mano.')
bullet(doc, 'Si tu área detecta una bandera (alergia nueva, condición médica, lazo que apareció), se reporta al Co-Dir y vuelve al tablero.')

espacio_notas(doc, 'Banderas / alertas del tablero hoy', 3)

# ---------- §6 RACI por equipo ----------
h1(doc, '6 · RACI por equipo (responsabilidades de cada pieza)')
para(doc, 'R = Responsable (hace) · A = Accountable (rinde cuentas) · C = Consultado · I = Informado')
make_table(doc,
    headers=['Pieza del retiro', 'Co-Dir', 'Asesores', 'Guías', 'Cocina', 'Música'],
    rows=[
        ['Testimonios', 'A', 'C', 'R', 'I', 'I'],
        ['Pequeños Grupos (PG)', 'A', 'C', 'R', 'I', 'I'],
        ['Dinámicas de PG', 'A', 'C', 'R', 'I', 'I'],
        ['Decoración del plenario', 'A', 'C', 'R', 'C', 'C'],
        ['Palancas (escritura)', 'A', 'C', 'R', 'C', 'C'],
        ['Entrega de palancas', 'A', 'R', 'C', 'C', 'C'],
        ['Carnets de PG', 'A', 'I', 'R', 'I', 'I'],
        ['Bolsita del participante', 'A', 'C', 'R', 'I', 'I'],
        ['Sociodramas', 'A', 'C', 'R', 'I', 'I'],
        ['Banderín', 'A', 'C', 'R', 'I', 'I'],
        ['Comida (desayuno · almuerzo · cena · refrigerios)', 'A', 'C', 'I', 'R', 'I'],
        ['Ambientación del comedor', 'A', 'C', 'I', 'R', 'I'],
        ['Lavatorio del sábado', 'A', 'C', 'C', 'R', 'C'],
        ['Sonido / instalación', 'A', 'C', 'I', 'I', 'R'],
        ['Música ambiente / despertar / transiciones', 'A', 'I', 'I', 'I', 'R'],
        ['Bayuyo (Sal y Luz del Mundo)', 'A', 'C', 'C', 'I', 'R'],
        ['Cancionero impreso', 'A', 'I', 'I', 'I', 'R'],
        ['Cantos de la Misa de Clausura', 'A', 'C', 'I', 'I', 'R'],
        ['Confesiones (coordinación)', 'A', 'R', 'I', 'I', 'C'],
        ['Eucaristía de Clausura', 'A', 'R', 'C', 'C', 'C'],
        ['Carta de compromiso (proyecto de vida)', 'A', 'R', 'C', 'I', 'I'],
        ['Entrega de los peces', 'R', 'C', 'I', 'I', 'I'],
        ['Avanzada en la casa', 'A', 'C', 'I', 'R', 'C'],
        ['Visita previa a la casa (con admin.)', 'R', 'C', 'C', 'C', 'I'],
        ['Captación de participantes', 'R', 'C', 'C', 'I', 'I'],
        ['Cobro de cuotas mensuales del equipo', 'A', 'C', 'R', 'R', 'R'],
    ],
    col_widths_cm=[7.5, 1.7, 1.7, 1.7, 1.7, 1.7])

espacio_notas(doc, 'Piezas que faltan o que cambian (escribir arriba)', 4)
doc.add_page_break()

# ---------- §7 MATERIALES DEL RETIRO ----------
h1(doc, '7 · Materiales del retiro (por área)')
para(doc, 'Esta es la lista de partida — se completa por área en sus propias reuniones. Lo de cada área queda en su anexo correspondiente.')

h2(doc, '7.1 Materiales generales / Co-Dirección')
bullet(doc, '**Biblias** (una por participante).')
bullet(doc, '**Peces de la clausura** (uno por participante).')
bullet(doc, '**Banderín** del ETC 88.')
bullet(doc, '**Hojas de datos personales** (para el 1er PG).')
bullet(doc, '**Cajas para palancas y biblias.**')
bullet(doc, '**Papelógrafo** (capilla, sábado tarde).')
bullet(doc, '**Marcadores, lapiceros, cinta adhesiva, tijeras.**')
bullet(doc, '**Sobres manila** (2do PG).')
bullet(doc, '**Cartulinas grandes** (dinámica de la tarjeta).')
bullet(doc, '**Material para reglas de la casa** (cartel o impresión).')

h2(doc, '7.2 Bolsita del participante')
para(doc, 'La bolsita se entrega al participante al inicio del retiro. Contenido propuesto:')
bullet(doc, '**Carnet del PG** (con su color).')
bullet(doc, '**Libreta** + **lapicero** (para notas durante el retiro).')
bullet(doc, '**Biblia** *(si va dentro o se entrega aparte en el momento de la entrega de la biblia — definir).*')
bullet(doc, '**Pez de cierre** (que se entrega en la Eucaristía final).')
bullet(doc, '**Hoja de datos / nombre.**')
linea_para_escribir(doc, 'Otros:')
linea_para_escribir(doc, 'Otros:')

h2(doc, '7.3 Materiales del equipo de Guías')
para(doc, '*Detalle por pequeño grupo en la Guía de Guías §4.*')
make_table(doc,
    headers=['Momento', 'Materiales'],
    rows=[
        ['1er PG (vie 7:35pm)', 'Hoja de datos · lápices'],
        ['2do PG (vie 8:30pm)', 'Sobre manila · marcadores · tarjetas · carnets'],
        ['3er PG (vie 10:40pm)', 'Libretas · lapiceros · **abrigo**'],
        ['4to PG (sáb 8:35am)', '**Alambre** · Biblia · hojas (si necesario)'],
        ['5to PG (sáb 10:45am)', '(preparación de sociodrama)'],
        ['6to PG (sáb 11:45am)', 'Objeto/símbolo (cada participante busca un objeto de la naturaleza)'],
        ['7mo PG (sáb 4:00pm)', '(ensayo del sociodrama)'],
        ['8vo PG (sáb 7:50pm)', '(comentar Regalo del Perdón)'],
        ['Último PG (dom 9:00am)', '**Venda** (dinámica del cariño)'],
        ['Por guía', 'Pañuelo · abrigo personal · palanca a cada participante'],
    ],
    col_widths_cm=[5, 12])

h2(doc, '7.4 Materiales del equipo de Cocina')
bullet(doc, '**Palangana** + **toallas** + **agua tibia** (Lavatorio del sábado).')
bullet(doc, '**Motivos por tiempo de comida** (decoración del comedor según el tema del día).')
bullet(doc, '**Utensilios** propios + complementos de la casa.')
bullet(doc, '**Material para servir a la mesa** (bandejas, servilleteros, etc.).')
bullet(doc, '**Bendiciones breves** por tiempo de comida (papel, ligado al tema).')
bullet(doc, '**Refrigerios** (AM, PM, noche).')
bullet(doc, '**Material de avanzada** (compras de no perecederos previas).')

h2(doc, '7.4b Menú del retiro (a llenar en vivo con Cocina)')
para(doc, 'Borrador del menú por tiempo de comida. Se cierra con Cocina **después de F1**, cuando se conoce la temática. La columna *Motivo* es la ambientación / mensaje del momento (sociodrama, oración, gesto). La cena del sábado lleva siempre el **Lavatorio (Jn 13)**.')
make_table(doc,
    headers=['Día', 'Tiempo', 'Hora', 'Menú', 'Motivo / Ambientación'],
    rows=[
        ['Viernes', 'Cena', '8:00 PM', '_____________', '_____________'],
        ['Sábado',  'Desayuno', '8:00 AM', '_____________', '_____________'],
        ['Sábado',  'Refrigerio AM', '10:25 AM', '_____________', '_____________'],
        ['Sábado',  'Almuerzo', '12:00 PM', '_____________', '_____________'],
        ['Sábado',  'Refrigerio PM', '4:45 PM', '_____________', '_____________'],
        ['Sábado',  'Cena con Lavatorio', '6:30 PM', '_____________', 'Jn 13 · Lavatorio de los pies'],
        ['Sábado',  'Refrigerio noche', 'post-Bayuyo', '_____________', '_____________'],
        ['Domingo', 'Desayuno', '8:00 AM', '_____________', '_____________'],
        ['Domingo', 'Almuerzo', '1:00 PM', '_____________', '_____________'],
    ],
    col_widths_cm=[2.0, 3.2, 2.2, 5.6, 4.0])
para(doc)
bullet(doc, '**Refrigerios** especiales: cumpleaños del retiro, momento de palancas generales, etc. — Cocina propone.')
bullet(doc, '**Donaciones en especie** (arroz, habichuelas, aceite) reducen ~15-25% el costo de la canasta — confirmar con Cocina cuáles se solicitan.')
bullet(doc, '**Alergias** del tablero a respetar al cerrar cada plato (mariscos · piña · huevo · canela · gastritis · diabetes · presión).')

h2(doc, '7.5 Materiales del equipo de Música')
bullet(doc, '**Sonido:** consola + parlantes (casa o respaldo del equipo).')
bullet(doc, '**Micrófonos** vocales + instrumentos · pies de mic · cables · extensiones.')
bullet(doc, '**Atriles** para partituras / tablets.')
bullet(doc, '**Instrumentos:** guitarras, cajón, teclado, etc.')
bullet(doc, '**Cancionero impreso** para participantes.')
bullet(doc, '**Cantos de la Misa de Clausura impresos.**')
bullet(doc, '**Antifaces de tela negra** (Bayuyo · uno por participante).')
bullet(doc, '**Playlist** preparada en un dispositivo (música ambiente, despertar, transiciones).')
bullet(doc, '**Setlists impresos** por momento del retiro.')

h2(doc, '7.6 Materiales litúrgicos')
bullet(doc, '**Vino litúrgico** (la parroquia del Padre Paul lo aporta).')
bullet(doc, '**Hostias.**')
bullet(doc, '**Ornamento** del sacerdote.')
bullet(doc, '**Vasos sagrados** (cáliz, patena, copón).')
bullet(doc, '**Biblias** (entrega del sábado tarde).')
bullet(doc, '**Banderín** (sábado noche · preparación).')
bullet(doc, '**Velas** para confesiones / Lavatorio.')

h2(doc, '7.7 Materiales para el Rosario / Oración a María (sábado mañana)')
bullet(doc, '**Rosarios** (uno por participante · pueden ir en la bolsita).')
bullet(doc, '**Cantos marianos impresos** (Música).')
bullet(doc, '**Imagen mariana** (si la casa no la tiene).')

espacio_notas(doc, 'Materiales que faltan o que cambian (escribir arriba)', 5)
doc.add_page_break()

# ---------- §8 PAGOS MENSUALES ----------
h1(doc, '8 · Pagos mensuales del equipo')
para(doc, 'Cada miembro del equipo aporta una **cuota mensual** que cubre su participación en el retiro y los gastos comunes del equipo.')

h3(doc, 'Monto y calendario')
linea_para_escribir(doc, 'Monto mensual:')
linea_para_escribir(doc, 'Fecha de cobro:')
linea_para_escribir(doc, 'Cantidad de meses:')
linea_para_escribir(doc, 'Qué cubre:')

h3(doc, 'Cómo se cobra')
bullet(doc, 'Cada **coord cobra a su equipo** y reporta lo cobrado al Co-Dir mensualmente (o a Tesorería cuando esté nombrada).')
bullet(doc, 'Se registra en la hoja de **Control de Pagos** del tablero.')
bullet(doc, 'El Co-Dir o Tesorería confirma cada mes que las cuotas están al día.')

h3(doc, 'Si alguien no puede pagar')
bullet(doc, 'El coord lo escala al Co-Dir — **nadie queda fuera por dinero**.')
bullet(doc, 'Hay **padrinazgo** (un padrino cubre la cuota) y **arreglos posibles** (pago escalonado, descuento).')
bullet(doc, 'Se conversa de manera reservada; el caso no se comparte con el resto del equipo.')

espacio_notas(doc, 'Cuota mensual a confirmar en vivo', 4)
doc.add_page_break()

# ---------- §9 REQUERIMIENTOS POR EQUIPO + ROADMAP ----------
h1(doc, '9 · Qué requerimos de cada equipo + Roadmap')

h2(doc, 'Guías — Priscilla + Camila')
para(doc, '**Para F1:**')
bullet(doc, 'Presupuesto de Guías (libretas de los participantes entran aquí).')
bullet(doc, 'Lista de materiales del área.')
bullet(doc, '3 dudas o brechas del equipo.')
para(doc, '*Cómo funciona el área:* ver **Guía de Guías 88**.')

h2(doc, 'Cocina — Paloma + Jhonnito')
para(doc, '**Para F1:**')
bullet(doc, '**Primero borrador de menú**, luego presupuesto base.')
bullet(doc, '3 dudas o brechas del equipo.')
para(doc, '*Cómo funciona el área:* ver **Anexo de Cocina 88**.')

h2(doc, 'Música — José Ángel Tusen')
para(doc, '**Para F1:**')
bullet(doc, 'Presupuesto base.')
bullet(doc, 'Repertorio borrador y cancionero base.')
bullet(doc, '3 dudas o brechas del equipo.')
para(doc, '*Cómo funciona el área:* ver **Anexo de Música 88**.')

h2(doc, 'Calendario de preparación')
make_table(doc,
    headers=['Hito', 'Cuándo', 'Foco'],
    rows=[
        ['Reunión de coords', 'jueves 4-jun (virtual)', 'Esta reunión'],
        ['F1 · 1ra Formación', '[PROPUESTA] 14-jun', 'Conformación · cantera · ¿qué es un ETC? · Carpeta'],
        ['F2 · 2da Formación', '[PROPUESTA] 28-jun', 'Tema 1 · primeros testimonios'],
        ['F3 · 3ra Formación', '5-jul', 'Tema 2 · inicio recaudación'],
        ['F4 · 4ta Formación', '[PROPUESTA] 19-jul', 'Tema 3 · perfiles 1 y 2'],
        ['F5 · 5ta Formación', '[PROPUESTA] 2-ago', 'Tema 4 · cierre logístico'],
        ['Profondo', '31-jul → 2-ago', 'Actividad pro-fondos del equipo'],
        ['Visita a la casa', '[POR DEFINIR]', 'Cocina + Guías + Co-Dir + admin (Samuel Montilla)'],
        ['Convivencia + mini-retiro', '22-ago', 'Día completo'],
        ['Ensayo General', '23-ago (obligatorio)', 'Pase completo del retiro'],
        ['Avanzada', '3-sep', 'Cocina + Música anticipados'],
        ['ETC 88', '4–6 sep 2026', 'Higüey · La Ceiba del Salado'],
    ],
    col_widths_cm=[5, 4, 8])

espacio_notas(doc, 'Cambios al roadmap en vivo', 4)
doc.add_page_break()

# ---------- §10 DISCIPLINA Y ASISTENCIA ----------
h1(doc, '10 · Disciplina y asistencia')

h2(doc, 'Asistencia')
bullet(doc, 'La base es **venir**. La presencia ya es señal de compromiso.')
bullet(doc, 'Inasistencias justificadas: **hasta 3 reuniones** (contando extraordinarias). Si se excede, conversa el coord con el Co-Dir.')
bullet(doc, '**Ensayo General (23-ago): asistencia obligatoria.** Sin ensayo, no se sirve en el retiro.')
bullet(doc, 'Cada coord lleva el registro semanal de asistencia de su área y lo comparte con el Co-Dir.')

h2(doc, 'Cómo escalar al Co-Dir (proceso recomendado)')
numbered(doc, 'El coord conversa primero con la persona en privado, sin juicio, buscando entender.')
numbered(doc, 'Si la situación se sostiene, el coord la trae al Co-Dir.')
numbered(doc, 'El Co-Dir **ora con los Asesores antes de cualquier decisión**, poniendo siempre a la persona por encima de la regla.')
numbered(doc, 'La decisión se comunica primero a la persona involucrada y después al equipo si corresponde.')

h2(doc, 'Conflictos entre miembros del equipo')
bullet(doc, 'El coord intenta el primer puente. Si el conflicto persiste, lo escala al Co-Dir + Asesores.')
bullet(doc, 'Confidencialidad estricta: lo del equipo se queda en el equipo.')

h2(doc, 'Renuncia o salida')
bullet(doc, 'El coord avisa al Co-Dir tan pronto sepa.')
bullet(doc, 'La sustitución se decide con los Asesores (cantera, antes que externos).')

h2(doc, 'Confidencialidad')
bullet(doc, 'Lo que se discute en reunión de equipo **no sale del equipo**.')
bullet(doc, '**Los testimonios de vida son confidenciales para siempre** (regla del guía).')
bullet(doc, 'Información médica / alergias / lazos del tablero: solo se usa para servir, no se comenta fuera.')

espacio_notas(doc, 'Casos que hay que hablar / conversar con los directores', 6)
doc.add_page_break()

# ---------- §11 OTROS TEMAS QUE PROPONGO ----------
h1(doc, '11 · Otros temas que propongo (revisar conmigo qué dejamos)')
para(doc, 'Estos temas no estaban en la lista original pero pueden agregarse o cubrirse en otra sesión. Cada uno con una nota de por qué.')

h3(doc, '11.1 · Captación de participantes')
para(doc, 'Cómo se busca, se evalúa y se ubica a los participantes (perfiles 1 y 2). Hoy va por Co-Dir y Asesores de comunidad, pero cada coord puede traer prospectos. ¿Lo trabajamos en F2/F3 o tiene espacio aquí?')
linea_para_escribir(doc, 'Decisión:')

h3(doc, '11.2 · Comunicación del equipo')
para(doc, 'Canales del equipo (grupo de WhatsApp del retiro · chat por área · canal de avisos del Co-Dir · dónde vive cada cosa). Definir un solo canal "oficial" para que no se diluya.')
linea_para_escribir(doc, 'Decisión:')

h3(doc, '11.3 · Visita previa a la casa')
para(doc, 'Cocina + Guías + Co-Dir + administración (Samuel Montilla). Define utensilios, espacio del PG, ambientación. **Es la pieza que más ha faltado en ediciones anteriores.**')
linea_para_escribir(doc, 'Fecha propuesta:')

h3(doc, '11.4 · Espiritualidad del equipo (oración + mini-retiro)')
para(doc, 'La Convivencia del 22-ago tiene mini-retiro. ¿Hay también momentos de oración del equipo entre formaciones? Vale para los coords con su gente.')
linea_para_escribir(doc, 'Decisión:')

h3(doc, '11.5 · Briefing pastoral del director')
para(doc, '¿Hay un mensaje del Co-Dir que cada coord debe transmitir al inicio de su 1ra reunión? Útil si quieres una sola voz desde el principio.')
linea_para_escribir(doc, 'Mensaje a transmitir:')

h3(doc, '11.6 · Plan de contingencias')
para(doc, '¿Qué hacemos si un guía se enferma 2 semanas antes? ¿Si la cocina no completa? ¿Si llueve fuerte el sábado? Una matriz corta.')
linea_para_escribir(doc, 'Decisión:')

h3(doc, '11.7 · Equipos auxiliares (los 5)')
para(doc, 'Recaudación · Guagua · Profondo · Intersección · Tesorería. ¿Salimos hoy con un primer nombre por equipo, o se trabaja en F1/F2?')
linea_para_escribir(doc, 'Decisión:')

h3(doc, '11.8 · Glosario eteciano para los nuevos')
para(doc, 'PG · Plenario · Palanca · Bayuyo · Lavatorio · Cuarto Día · Profondo · Intersección. Algunos miembros nuevos del equipo no los conocen. ¿Hacemos un anexo corto?')
linea_para_escribir(doc, 'Decisión:')

h3(doc, '11.9 · Recursos: dónde vive cada cosa')
para(doc, 'Link del tablero · link del repo · drive de entregables · drive de trabajo · agenda del director. Una sola hoja con todos los enlaces.')
linea_para_escribir(doc, 'Decisión:')

h3(doc, '11.10 · Foto del equipo + comunicación visual')
para(doc, 'Foto del equipo en F1 para usar en RRSS · plantillas de WhatsApp por momento del retiro · separadores de la Carpeta. Va a Claude Design.')
linea_para_escribir(doc, 'Decisión:')

espacio_notas(doc, 'Otros temas que el director quiere añadir', 5)
doc.add_page_break()

# ---------- §12 DECISIONES Y NOTAS EN VIVO ----------
h1(doc, '12 · Decisiones, acuerdos y notas en vivo')
para(doc, 'Espacio libre para llenar durante la reunión.')

h2(doc, 'Compromisos por coord para F1')
make_table(doc,
    headers=['Coord', 'Entregables F1', 'Fecha · Lugar próxima reunión'],
    rows=[
        ['Priscilla + Camila (Guías)', '_________________', '_________________'],
        ['Paloma + Jhonnito (Cocina)', '_________________', '_________________'],
        ['José Ángel (Música)', '_________________', '_________________'],
    ],
    col_widths_cm=[5, 6, 6])

h2(doc, 'Parqueo (dudas para responder por chat)')
for _ in range(6):
    linea_para_escribir(doc)

h2(doc, 'Próximos pasos del Co-Dir')
for _ in range(5):
    linea_para_escribir(doc)

# Pie
para(doc)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Documento de Trabajo · ETC 88 · 4-jun-2026 · todo lo del equipo es [PROPUESTA] hasta confirmación del Co-Dir.')
r.font.name = FUENTE; r.font.size = Pt(9); r.italic = True; r.font.color.rgb = GRIS

# ===== Guardar =====
out_path = os.path.join(OUT_DIR, 'DOC_TRABAJO_REUNION_COORDS.docx')
doc.save(out_path)
print(f"Wrote {out_path}")
