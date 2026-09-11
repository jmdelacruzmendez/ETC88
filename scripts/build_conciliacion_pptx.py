#!/usr/bin/env python3
"""
Presentación de la Conciliación Final ETC 88 (12 láminas, 16:9).
Se genera desde los MISMOS datos y cálculos de scripts/build_conciliacion.py: ejecuta su bloque de
datos y cálculos (hasta antes de escribir hojas) y arma las láminas. No se edita a mano (regla #2).
Salida: data/presupuesto/Conciliacion_ETC88_Presentacion.pptx
Requiere: pip install python-pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

REPO = Path(__file__).resolve().parent.parent
SRC = REPO / 'scripts/build_conciliacion.py'
OUT = REPO / 'data/presupuesto/Conciliacion_ETC88_Presentacion.pptx'

# ── datos y cálculos compartidos con el Excel ──
src = SRC.read_text(encoding='utf-8')
MARK = '# ══════════ HELPERS DE HOJA ══════════'
assert MARK in src, 'build_conciliacion.py cambió de estructura'
ns = {'__file__': str(SRC)}
exec(compile(src.split(MARK)[0], str(SRC), 'exec'), ns)
D = ns  # acceso corto

MAR = '5B3A29'; AZUL = '0B1F3A'; VERDE = '10B981'; ROJO = 'EF4444'; CREMA = 'F7EFD9'; AMBAR = 'F59E0B'; GRIS = '888888'; TINTA = '1E293B'; ARENA = 'F7F3EA'
def rgb(h): return RGBColor.from_string(h)
def fmt(v):
    if isinstance(v, bool): return str(v)
    if isinstance(v, float): return f'{v:,.2f}' if abs(v - round(v)) > 0.005 else f'{v:,.0f}'
    if isinstance(v, int): return f'{v:,}'
    return str(v)

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
FOOT = 'ETC 88 · Conciliación final · cierre 11-sep-2026 · generado desde datos (scripts/build_conciliacion_pptx.py)'

def text(s, left, top, width, height, runs, size=14, color=TINTA, bold=False, align=PP_ALIGN.LEFT, italic=False):
    tb = s.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = runs; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = rgb(color)
    return tb

def new_slide(title, kicker=None, color=AZUL):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.1)); bar.fill.solid(); bar.fill.fore_color.rgb = rgb(color); bar.line.fill.background()
    text(s, Inches(0.5), Inches(0.22), W - Inches(1), Inches(0.7), title, size=26, color=CREMA, bold=True)
    if kicker: text(s, Inches(0.5), Inches(1.2), W - Inches(1), Inches(0.5), kicker, size=13, color=GRIS, italic=True)
    text(s, Inches(0.5), H - Inches(0.42), W - Inches(1), Inches(0.3), FOOT, size=9, color=GRIS)
    return s

def bullets(s, items, left, top, width, height, size=15, color=TINTA, gap=6):
    tb = s.shapes.add_textbox(left, top, width, height); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap)
        if isinstance(it, tuple):
            head, rest = it
            r1 = p.add_run(); r1.text = head; r1.font.bold = True; r1.font.size = Pt(size); r1.font.color.rgb = rgb(color)
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size); r2.font.color.rgb = rgb(color)
        else:
            r = p.add_run(); r.text = it; r.font.size = Pt(size); r.font.color.rgb = rgb(color)
    return tb

def table(s, rows, left, top, width, col_w=None, size=12, row_h=Inches(0.36), bold_last=False):
    nrows, ncols = len(rows), len(rows[0])
    shp = s.shapes.add_table(nrows, ncols, left, top, width, row_h * nrows); t = shp.table
    if col_w:
        for i, w in enumerate(col_w): t.columns[i].width = w
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
            r = p.add_run(); r.text = fmt(v); r.font.size = Pt(size)
            if isinstance(v, (int, float)): p.alignment = PP_ALIGN.RIGHT
            c.margin_left = c.margin_right = Inches(0.07); c.margin_top = c.margin_bottom = Inches(0.03)
            if i == 0:
                r.font.bold = True; r.font.color.rgb = rgb(CREMA); c.fill.solid(); c.fill.fore_color.rgb = rgb(MAR)
            else:
                c.fill.solid(); c.fill.fore_color.rgb = rgb('FFFFFF' if i % 2 else ARENA); r.font.color.rgb = rgb(TINTA)
                if bold_last and i == nrows - 1: r.font.bold = True; c.fill.fore_color.rgb = rgb('E3F2FD')
    return t

def cards(s, items, top, height=Inches(1.55), left0=Inches(0.5), total_w=None):
    n = len(items); gap = Inches(0.22); total_w = total_w or (W - Inches(1)); w = int((total_w - gap * (n - 1)) / n)
    for i, (label, val, sub, color) in enumerate(items):
        left = left0 + (w + gap) * i
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, height); box.fill.solid(); box.fill.fore_color.rgb = rgb(ARENA); box.line.color.rgb = rgb(color); box.line.width = Pt(1.5)
        tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = label.upper(); r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = rgb(GRIS)
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; r2 = p2.add_run(); r2.text = val; r2.font.size = Pt(22); r2.font.bold = True; r2.font.color.rgb = rgb(color)
        p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER; r3 = p3.add_run(); r3.text = sub; r3.font.size = Pt(9); r3.font.color.rgb = rgb(TINTA)

def bar_chart(s, cats, series, left, top, width, height, stacked=False, title=None, colors=(MAR, VERDE, AZUL, AMBAR, ROJO), labels=True):
    cd = CategoryChartData(); cd.categories = cats
    for name, vals in series: cd.add_series(name, vals)
    ct = XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED
    ch = s.shapes.add_chart(ct, left, top, width, height, cd).chart
    ch.has_legend = len(series) > 1
    if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(10)
    plot = ch.plots[0]; plot.has_data_labels = labels
    if labels: plot.data_labels.number_format = '#,##0'; plot.data_labels.number_format_is_linked = False; plot.data_labels.font.size = Pt(9)
    plot.gap_width = 60
    ch.category_axis.tick_labels.font.size = Pt(10); ch.value_axis.tick_labels.font.size = Pt(9); ch.value_axis.has_major_gridlines = False
    ch.value_axis.tick_labels.number_format = '#,##0'; ch.value_axis.tick_labels.number_format_is_linked = False
    for i, ser in enumerate(ch.series): f = ser.format.fill; f.solid(); f.fore_color.rgb = rgb(colors[i % len(colors)])
    if title: ch.has_title = True; ch.chart_title.text_frame.text = title; ch.chart_title.text_frame.paragraphs[0].font.size = Pt(12); ch.chart_title.text_frame.paragraphs[0].font.bold = True
    else: ch.has_title = False
    return ch

# ── cifras ──
ENTRADAS, SALIDAS, BALANCE = D['ENTRADAS'], D['SALIDAS'], D['BALANCE']
REAL, IMPUESTOS = D['REAL_CUENTAS'], D['IMPUESTOS']
COSTO_ECON, GRATIS = D['COSTO_ECON'], D['RECIBIDO_GRATIS']; COSTO_COMPLETO = COSTO_ECON
ESPECIE, CORT = D['ESPECIE_TOTAL'], D['CORTESIA_JUEVES'] + D['CORTESIA_HOSP']
PPTO, PPTO_S, OFICIAL = D['PPTO_11AGO'], D['PPTO_SISTEM'], D['PPTO_OFICIAL']
PERSONAS, FUENTES, AREAS, cruce = D['PERSONAS'], D['FUENTES'], D['AREAS'], D['cruce']
pendientes, checks, AUDIT_PASS, AUDIT_N = D['pendientes'], D['checks'], D['AUDIT_PASS'], D['AUDIT_N']
FORM = 0; HAB_T, HAB_TAR = D['HAB_TOTAL'], D['HAB_TARIFA']; CASA_TOTAL = D['CASA_TOTAL']
tipo_89 = D['tipo_89']
PUNTUALES = D['PUNTUALES']; BASE_REC = D['BASE_REC']
abiertos = [x for x in pendientes if x[0] == 'ABIERTO']; notas = [x for x in pendientes if x[0] == 'NOTA']

# ── 1. Portada ──
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H); bg.fill.solid(); bg.fill.fore_color.rgb = rgb(AZUL); bg.line.fill.background()
text(s, Inches(0.8), Inches(1.6), W - Inches(1.6), Inches(1.2), 'ETC 88 · Conciliación final', size=44, color=CREMA, bold=True)
text(s, Inches(0.8), Inches(2.8), W - Inches(1.6), Inches(0.6), 'Encuentro Total con Cristo #88 · Casa La Ceiba del Salado, Higüey · 4–6 de septiembre de 2026', size=18, color=CREMA)
text(s, Inches(0.8), Inches(3.5), W - Inches(1.6), Inches(0.6), 'Co-Dirección: Juan Manuel de la Cruz · Jean Carlo de la Cruz', size=16, color=CREMA)
text(s, Inches(0.8), Inches(4.6), W - Inches(1.6), Inches(0.5), f'Cierre al 11-sep-2026 · presupuesto oficial: {OFICIAL} ({PPTO:,.2f}) · {PERSONAS} personas en la casa', size=14, color=GRIS)
text(s, Inches(0.8), Inches(5.1), W - Inches(1.6), Inches(0.5), f'Cifras en RD$ · {AUDIT_N} verificaciones numéricas recalculadas ({AUDIT_PASS} PASS) · acompaña al Excel Conciliacion_Final_ETC88.xlsx', size=12, color=GRIS)

# ── 2. En una frase ──
s = new_slide('En una frase', 'lo que el Consejo necesita saber antes de abrir el Excel')
cards(s, [('Lo que costó el retiro', f'{COSTO_ECON:,.0f}', 'caja + especie + cortesías', AZUL),
          ('Lo que salió de caja', f'{SALIDAS:,}', '22 gastos', ROJO),
          ('Lo que quedó en cuentas', f'{REAL:,}', f'balance {BALANCE:,.2f} − impuestos {IMPUESTOS:,.2f}', VERDE)], Inches(1.8))
bullets(s, [
    ('El retiro costó más de lo presupuestado y salió menos de caja. ', f'Costó {COSTO_ECON:,.0f} ({(COSTO_ECON / PPTO - 1) * 100:+.1f}% sobre el presupuesto oficial); de caja, {SALIDAS:,} ({(SALIDAS / PPTO - 1) * 100:+.1f}%).'),
    ('La diferencia la pusieron otros. ', f'{GRATIS:,.0f} ({GRATIS / COSTO_ECON * 100:.0f}% del costo) llegaron sin pasar por caja: 30 donantes en especie y la casa en cortesías.'),
    ('Las cuentas cuadran. ', f'Entradas {ENTRADAS:,.2f} − salidas {SALIDAS:,} = {BALANCE:,.2f}; en el banco hay {REAL:,}; los {IMPUESTOS:,.2f} de diferencia son impuestos y comisiones.'),
    ('Cada persona costó ', f'{SALIDAS / PERSONAS:,.0f} de caja y {COSTO_ECON / PERSONAS:,.0f} al costo; el participante pagó 3,500.'),
], Inches(0.6), Inches(3.7), W - Inches(1.2), Inches(3.2), size=16, gap=10)

# ── 3. Cómo se hizo ──
s = new_slide('Cómo se hizo esta conciliación', 'el proceso, paso a paso')
bullets(s, [
    ('1 · Fuentes. ', 'Registro de tesorería del tablero de finanzas (22 gastos, 49 donaciones, 8 entregas del profondo), totales del tablero, estado de cuenta al 11-sep, los dos presupuestos (11-Ago y Sistem), la cotización de Iberia, las listas y aclaraciones del director (10 y 11 sep) y la caja de junio del Presupuesto Maestro.'),
    ('2 · Una sola caja. ', 'El registro de tesorería es la fuente única de lo que entró y salió. Cuadra al peso con el tablero y, tras los impuestos, con el banco.'),
    ('3 · Cruce con el presupuesto. ', f'Cada gasto se asignó a una partida del presupuesto oficial ({OFICIAL}); el presupuesto por partida cuadra al centavo con el archivo. El pago único a JM de 37,500 se abrió en sus tres conceptos.'),
    ('4 · Lo que no pasó por caja. ', 'Las donaciones en especie se valoraron a precio de presupuesto (no a factura) y las cortesías de la casa como valor recibido.'),
    ('5 · Lo que costó. ', f'caja {SALIDAS:,} + especie {ESPECIE:,} + cortesías {CORT:,} = {COSTO_ECON:,.0f}.'),
    ('6 · Auditoría. ', f'{AUDIT_N} verificaciones se recalculan cada vez que se genera el archivo; {AUDIT_PASS} pasan y la restante es un dato pendiente, no un error de suma.'),
    ('7 · Nada a mano. ', 'Excel y presentación salen de un mismo script a partir de los datos; para corregir algo se cambia el dato y se regenera.'),
], Inches(0.6), Inches(1.8), W - Inches(1.2), Inches(5.2), size=14, gap=9)

# ── 4. Los seis números ──
s = new_slide('Los seis números', 'todo lo demás es detalle de estos')
cards(s, [('Entradas', f'{ENTRADAS:,.2f}', 'participantes + donaciones + profondo + cuotas + tardanzas', MAR),
          ('Salidas de caja', f'{SALIDAS:,}', '22 gastos', ROJO),
          ('Balance → banco', f'{BALANCE:,.2f}', f'{REAL:,} en cuentas · {IMPUESTOS:,.2f} impuestos', AZUL)], Inches(1.8))
cards(s, [('Costo económico', f'{COSTO_ECON:,.0f}', f'caja + especie ({ESPECIE:,}) + cortesías ({CORT:,})', AZUL),
          ('Recibido sin pagar', f'{GRATIS:,.0f}', f'{GRATIS / COSTO_ECON * 100:.1f}% del costo económico', VERDE),
          ('Por persona (99)', f'{SALIDAS / PERSONAS:,.0f} / {COSTO_ECON / PERSONAS:,.0f}', 'de caja / al costo · cuota del participante 3,500', MAR)], Inches(3.7))
text(s, Inches(0.6), Inches(5.6), W - Inches(1.2), Inches(1), f'Base para presupuestar el ETC 89: {COSTO_ECON:,.0f} en total, {COSTO_ECON / PERSONAS:,.0f} por persona.', size=13, color=GRIS, italic=True)

# ── 5. De dónde salió el dinero ──
s = new_slide('De dónde salió el dinero', f'entradas {ENTRADAS:,.2f} · cinco fuentes')
cats = [l.split(' (')[0].split(',')[0] for l, _ in FUENTES]; vals = [v for _, v in FUENTES]
bar_chart(s, cats, [('RD$', vals)], Inches(0.5), Inches(1.7), Inches(7.3), Inches(5.1))
rows = [['Fuente', 'RD$', '%']] + [[c, v, f'{v / ENTRADAS * 100:.1f}%'] for c, v in zip(cats, vals)] + [['Entradas totales', ENTRADAS, '100%']]
table(s, rows, Inches(8.1), Inches(1.9), Inches(4.7), [Inches(2.6), Inches(1.3), Inches(0.8)], size=12, bold_last=True)
text(s, Inches(8.1), Inches(4.7), Inches(4.7), Inches(1.6), 'Participantes y donaciones aportaron por igual; el profondo (rifa y venta de helados, en neto) fue la tercera fuente. Las cuotas del equipo cubrieron menos de una quinta parte.', size=12, color=GRIS, italic=True)

# ── 6. En qué se gastó ──
s = new_slide('En qué se gastó, por área', 'pagado de caja vs recibido sin pagar · costo real = pagado + donado')
rows = [['Área', f'Ppto oficial', 'Pagado', 'Donado', 'Costo real']]
for a, (p11, ps, pag, don) in AREAS.items():
    extra = FORM if a == 'Formación' else 0
    rows.append([a + (' (+ formaciones 10,000)' if extra else ''), p11, pag + extra, don, pag + don + extra])
rows.append(['Total', PPTO, SALIDAS + FORM, D['CRUCE_DONADO'], COSTO_COMPLETO])
table(s, rows, Inches(0.5), Inches(1.7), Inches(7.6), [Inches(2.8), Inches(1.2), Inches(1.2), Inches(1.2), Inches(1.2)], size=11, row_h=Inches(0.33), bold_last=True)
bar_chart(s, list(AREAS.keys()), [('Pagado', [v[2] for v in AREAS.values()]), ('Donado', [v[3] for v in AREAS.values()])], Inches(8.3), Inches(1.6), Inches(4.7), Inches(5.3), stacked=True, colors=(ROJO, VERDE), labels=False)
text(s, Inches(0.5), Inches(6.35), Inches(7.6), Inches(0.6), 'Costo real = pagado de caja + cubierto sin pagar (donaciones en especie y cortesías de la casa).', size=10, color=GRIS, italic=True)

# ── 7. La casa ──
s = new_slide('La casa, al detalle', f'Casa La Ceiba del Salado · total pagado {CASA_TOTAL:,}')
rows = [['Concepto', 'RD$'],
        ['Hospedaje base: 97 personas × 2,360 (vie–dom)', 228920],
        ['Noche del jueves (avanzada): 19 personas × 500', 9500],
        [f'Habitaciones para pequeños grupos: {HAB_TAR:,.0f} × 2 noches (tarifa inferida)', HAB_T],
        ['= Total pagado (avance 23,600 + final 216,620)', CASA_TOTAL],
        ['Cortesía: 8 personas del jueves sin cobro (8 × 500)', D['CORTESIA_JUEVES']],
        ['Cortesía: el padre y la sor, vie–dom (2 × 2,360)', D['CORTESIA_HOSP']],
        [f'vs presupuesto oficial (236,000)', CASA_TOTAL - 236000],
        ['vs presupuesto Sistem (200,000)', CASA_TOTAL - 200000]]
table(s, rows, Inches(0.5), Inches(1.7), Inches(7.4), [Inches(6.0), Inches(1.4)], size=12)
bullets(s, [
    ('Fuimos 99 y se facturaron 97. ', 'La casa dio dos cortesías: el padre y la sor (el director cedió la suya a la sor).'),
    ('Llegamos jueves, no viernes. ', '27 personas en la avanzada: 19 pagaron 500 y 8 no se cobraron. La comida de esa noche salió de la compra general.'),
    ('El Sistem quedó corto. ', 'Nunca actualizó la tarifa a 2,360 ni contempló el jueves; por eso el desvío de 40,220 contra ese presupuesto es aparente.'),
    ('Contra el oficial ', f'la casa costó {CASA_TOTAL - 236000:,} más: el jueves y las habitaciones.'),
], Inches(8.2), Inches(1.7), Inches(4.7), Inches(5), size=13, gap=8)

# ── 8. Presupuesto vs realidad ──
s = new_slide('Presupuesto oficial vs realidad', f'{OFICIAL}: {PPTO:,.2f} · costo completo con tarifa real de la casa')
cards(s, [(f'Presupuesto oficial', f'{PPTO:,.0f}', 'lo que se planificó (costo completo)', MAR),
          ('Costo económico real', f'{COSTO_ECON:,.0f}', f'{COSTO_ECON - PPTO:+,.0f} ({(COSTO_ECON / PPTO - 1) * 100:+.1f}%)', AZUL),
          ('Salidas de caja', f'{SALIDAS:,}', f'{SALIDAS - PPTO:+,.0f} ({(SALIDAS / PPTO - 1) * 100:+.1f}%)', VERDE)], Inches(1.8))
esp = D['especie_tot']
bullets(s, [
    ('Por qué valió más: ', f'la casa (+{CASA_TOTAL - 236000:,}: jueves y habitaciones), el bizcocho de bienvenida (6,000, no presupuestado), los helados del premio, las biblias (+1,360) y 15,000 de efectivo para imprevistos sin liquidar; a cambio, cocina y ofrendas quedaron por debajo.'),
    ('Por qué costó menos: ', f'{ESPECIE:,} en especie (peces 36,000 de La Vega; guías {esp["Guías"]:,}; cocina y decoración {esp["Cocina"]:,}; insumos de misa 3,000; oficina 2,300; música {esp["Música"]:,}) y {CORT:,} en cortesías de la casa.'),
    ('Lectura para el Consejo: ', 'el presupuesto oficial fue realista; lo que cambió fue quién pagó cada cosa.'),
], Inches(0.6), Inches(3.7), W - Inches(1.2), Inches(3.2), size=14, gap=10)

# ── 9. Balance y banco ──
s = new_slide('Balance y banco', 'las cuentas cuadran')
rows = [['Concepto', 'RD$', 'Nota'],
        ['Entradas totales', ENTRADAS, 'cinco fuentes'],
        ['Salidas de caja', SALIDAS, '22 gastos'],
        ['Balance', BALANCE, 'entradas − salidas'],
        ['Real en cuentas al 11-sep', REAL, 'cuenta del director + cuenta de Day'],
        ['Diferencia', IMPUESTOS, 'impuestos y comisiones bancarias'],
        ]
table(s, rows, Inches(0.5), Inches(1.7), Inches(12.3), [Inches(4.2), Inches(1.8), Inches(6.3)], size=12)
bullets(s, [
    ('Apoyo total recibido: ', f'{D["DONAC_EFEC"]:,} en efectivo + {ESPECIE:,} en especie + {CORT:,} en cortesías = {D["DONAC_EFEC"] + ESPECIE + CORT:,}.'),
    ('Qué hacer con los ', f'{REAL:,} es decisión de la dirección y del Consejo; esta conciliación no lo asume.'),
], Inches(0.6), Inches(4.6), W - Inches(1.2), Inches(2), size=14, gap=8)

# ── 10. Pendientes ──
s = new_slide('Pendientes y decisiones', f'{len(abiertos)} abiertos · {len(notas)} notas · el detalle completo está en la hoja 5 del Excel', color=AMBAR)
rows = [['Estado', 'Tema', 'Qué falta']] + [[e, t, f or d] for e, t, d, f in abiertos] + [[e, t, f or d] for e, t, d, f in notas]
table(s, rows, Inches(0.5), Inches(1.7), Inches(12.3), [Inches(1.1), Inches(3.2), Inches(8.0)], size=11, row_h=Inches(0.5))

# ── 11. Base ETC 89 ──
s = new_slide('Base para el ETC 89', 'el número real sobre el cual presupuestar', color=VERDE)
cards(s, [('Costo real 88', f'{COSTO_ECON:,.0f}', f'{COSTO_ECON / PERSONAS:,.0f} por persona · con todo lo donado valorado', AZUL),
          ('Base recurrente', f'{BASE_REC:,.0f}', f'{BASE_REC / PERSONAS:,.0f} por persona · sin desvío ni bizcocho de bienvenida', VERDE),
          ('Lo que hubo que financiar', f'{SALIDAS:,}', 'si el 89 no consigue las mismas donaciones, la cifra sube al costo completo', MAR)], Inches(1.8))
bullets(s, [
    ('Cómo usar la hoja 4 del Excel: ', 'cada partida trae su costo real y su tipo. "Por persona" se multiplica por asistentes (casa 2,360; comida ~1,061; biblias 680; peces 600; camisetas 480); "fijo" se cotiza; "puntual" se decide.'),
    ('Financiamiento del 88 como referencia: ', ' · '.join(f'{l.split(" (")[0].split(",")[0]} {v / ENTRADAS * 100:.0f}%' for l, v in FUENTES) + '.'),
    ('Lecciones: ', 'cerrar el presupuesto formalmente (los dos quedaron "Tentativo"); registrar cada donación en especie con su valor al recibirla; registro de tesorería con las mismas categorías que el presupuesto; liquidar los efectivos en 7 días; una línea por concepto.'),
], Inches(0.6), Inches(3.7), W - Inches(1.2), Inches(3.3), size=14, gap=10)

# ── 12. Cómo leer el Excel ──
s = new_slide('Cómo leer el Excel', 'Conciliacion_Final_ETC88.xlsx · 13 hojas, numeradas en el orden en que conviene leerlas', color=GRIS)
rows = [['Hoja', 'Qué responde'],
        ['0 Resumen', 'los seis números, tres respuestas, pendientes abiertos e índice'],
        ['1 Caja', 'entradas por fuente, salidas, balance, banco, apoyo total donado'],
        ['2 Cruce por partida', 'presupuesto oficial y Sistem → pagado → donado → costo real, por área (filas agrupadas)'],
        ['3 Al costo vs caja', 'qué vale el retiro y qué nos costó; por persona; la casa al detalle'],
        ['4 Base ETC 89', 'costo real por partida y por persona, con el tipo de cada partida para presupuestar'],
        ['5 Pendientes', 'una sola lista: abiertos, notas, cerrados e info, con qué falta en cada uno'],
        ['6 Ppto oficial vs Sistem', 'los dos presupuestos por área y su desvío'],
        ['7 Gastos · 8 Donaciones efectivo · 9 Especie · 10 Profondo', 'los anexos: cada movimiento con su nota'],
        ['11 Auditoría', f'{AUDIT_N} verificaciones recalculadas al generar el archivo'],
        ['12 Fuentes y método', 'de dónde sale cada número y las reglas usadas']]
table(s, rows, Inches(0.5), Inches(1.7), Inches(12.3), [Inches(4.3), Inches(8.0)], size=12, row_h=Inches(0.42))
text(s, Inches(0.5), Inches(6.5), Inches(12.3), Inches(0.5), 'Cada hoja lleva en su segunda fila una línea "Cómo leer" con lo que significa cada columna.', size=12, color=GRIS, italic=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f'✓ Wrote {OUT} ({len(prs.slides)} láminas)')
print(f'  costo económico {COSTO_ECON:,.0f} · completo {COSTO_COMPLETO:,.0f} · base recurrente {BASE_REC:,.0f} · abiertos {len(abiertos)} · auditoría {AUDIT_PASS}/{AUDIT_N}')
