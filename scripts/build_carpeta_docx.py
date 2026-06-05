#!/usr/bin/env python3
"""Renderiza la Carpeta del ETC 88 como un LIBRITO .docx: portada + índice +
las 12 secciones de la Carpeta + los 4 anexos (cada uno con su portadilla),
con numeración de página y pie. Estilo Dirección A. Es un borrador FUNCIONAL
para revisar contenido/estructura — el acabado visual final lo hace Claude Design.
Salida: entrega_diseno/CARPETA_ETC88.docx
"""
import os, re
import cairosvg
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CARPETA = os.path.join(REPO, 'preparacion', 'COPY_GUIA_ETC88.md')
ANEXOS = ['GUIA_DE_GUIAS_88.md', 'ANEXO_COCINA_88.md', 'ANEXO_MUSICA_88.md', 'ANEXO_ORACIONES_88.md']
OUT = os.path.join(REPO, 'entrega_diseno', 'CARPETA_ETC88.docx')
LOGO_SVG = os.path.join(REPO, 'design', 'logo_eteciano.svg')
LOGO_PNG = os.path.join(REPO, 'design', 'logo_eteciano.png')
os.makedirs(os.path.dirname(OUT), exist_ok=True)
if not os.path.exists(LOGO_PNG):
    cairosvg.svg2png(url=LOGO_SVG, write_to=LOGO_PNG, output_width=900)

AZUL, ARENA, CORAL, TINTA, GRIS = (RGBColor(0x1B,0x3A,0x5C), RGBColor(0xC0,0xA0,0x6C),
    RGBColor(0xE3,0x6C,0x4F), RGBColor(0x1A,0x1A,0x1A), RGBColor(0x66,0x66,0x66))
FUENTE = 'Calibri'

def set_cell_bg(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexc)
    tcPr.append(shd)

def add_runs(p, text, size=11, color=TINTA):
    for part in re.split(r'(\*\*[^*]+?\*\*|\*[^*\n]+?\*|`[^`]+?`)', text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            r = p.add_run(part[1:-1]); r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1])
        else:
            r = p.add_run(part)
        r.font.name = FUENTE; r.font.size = Pt(size); r.font.color.rgb = color

def h1(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.name = FUENTE; r.font.size = Pt(20); r.font.color.rgb = AZUL
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'12'); b.set(qn('w:space'),'4'); b.set(qn('w:color'),'E36C4F')
    pBdr.append(b); pPr.append(pBdr)

def h2(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.name = FUENTE; r.font.size = Pt(13); r.font.color.rgb = AZUL

def para(doc, text, italic=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(4)
    add_runs(p, text)
    if italic:
        for r in p.runs: r.italic = True; r.font.color.rgb = GRIS
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1)
    add_runs(p, text)

def numbered(doc, num, text):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f'{num}. '); r.bold = True; r.font.name = FUENTE; r.font.size = Pt(11); r.font.color.rgb = AZUL
    add_runs(p, text)

def callout(doc, lines):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0,0); set_cell_bg(cell, 'F7F1E3'); cell.width = Cm(16)
    tcPr = cell._tc.get_or_add_tcPr(); tcB = OxmlElement('w:tcBorders')
    for side,color,sz in [('left','E36C4F','24'),('top','E8D4A8','4'),('right','E8D4A8','4'),('bottom','E8D4A8','4')]:
        b = OxmlElement(f'w:{side}'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),sz); b.set(qn('w:color'),color); tcB.append(b)
    tcPr.append(tcB)
    cell.text = ''
    for j, ln in enumerate(lines):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1); add_runs(p, ln, size=10.5, color=AZUL)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def make_table(doc, headers, rows):
    t = doc.add_table(rows=len(rows)+1, cols=len(headers)); t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(headers):
        c = t.cell(0,j); set_cell_bg(c,'1B3A5C'); c.text=''
        rr = c.paragraphs[0].add_run(h); rr.bold=True; rr.font.name=FUENTE; rr.font.size=Pt(10); rr.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
    for i,rowv in enumerate(rows, start=1):
        for j,val in enumerate(rowv):
            c = t.cell(i,j); c.text=''; add_runs(c.paragraphs[0], val, size=10)

def page_field(paragraph):
    run = paragraph.add_run()
    for t, attr in [('begin', None), (None, 'PAGE'), ('end', None)]:
        if attr:
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = attr
        else:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), t)
        run._r.append(el)

def setup_footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True   # sin pie en la portada
    f = sec.footer
    p = f.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('ETC 88 · Higüey 2026   ·   '); r.font.name = FUENTE; r.font.size = Pt(8); r.font.color.rgb = GRIS
    page_field(p)
    for rr in p.runs: rr.font.size = Pt(8); rr.font.color.rgb = GRIS

def parse_body(doc, lines, page_break_per_section):
    """Renderiza el cuerpo (## ### #### > - 1. | texto). Sin lógica de portada."""
    i, first = 0, True
    while i < len(lines):
        line = lines[i].rstrip(); st = line.lstrip()
        if st.startswith('|'):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                block.append(lines[i].strip()); i += 1
            cells = lambda r: [c.strip() for c in r.strip().strip('|').split('|')]
            make_table(doc, cells(block[0]), [cells(b) for b in block[2:]] if len(block) > 2 else [])
            continue
        if st.startswith('>'):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith('>'):
                block.append(re.sub(r'^\s*>\s?', '', lines[i])); i += 1
            block = [b for b in block if b.strip()]
            if block: callout(doc, block)
            continue
        if line.startswith('## '):
            if page_break_per_section and not first: doc.add_page_break()
            first = False; h1(doc, line[3:]); i += 1; continue
        if line.startswith('### '):
            h2(doc, re.sub(r'\*', '', line[4:])); i += 1; continue
        if line.startswith('#### '):
            h2(doc, line[5:]); i += 1; continue
        if line.startswith('---') or line.strip() == '':
            i += 1; continue
        m = re.match(r'^(\d+)\.\s+(.*)', line)
        if m: numbered(doc, m.group(1), m.group(2)); i += 1; continue
        if st.startswith('- ') or st.startswith('* '):
            bullet(doc, st[2:]); i += 1; continue
        para(doc, line.strip()[1:-1], italic=True) if (line.strip().startswith('*') and line.strip().endswith('*') and line.count('*') == 2) else para(doc, line)
        i += 1

def split_cover(lines):
    """Devuelve (cover_lines, body_lines) partiendo en el primer '## '."""
    for idx, ln in enumerate(lines):
        if ln.startswith('## '):
            return lines[:idx], lines[idx:]
    return lines, []

def render_cover(doc, cover_lines):
    title = next((l[2:] for l in cover_lines if l.startswith('# ')), 'GUÍA DEL ETC 88')
    subs = [l[4:] for l in cover_lines if l.startswith('### ')]
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(LOGO_PNG, width=Inches(1.8)); p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold=True; r.font.name=FUENTE; r.font.size=Pt(30); r.font.color.rgb=AZUL
    for s in subs:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, s, size=14, color=ARENA)
        for r in p.runs: r.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    # callouts de portada
    i = 0
    while i < len(cover_lines):
        if cover_lines[i].lstrip().startswith('>'):
            block = []
            while i < len(cover_lines) and cover_lines[i].lstrip().startswith('>'):
                block.append(re.sub(r'^\s*>\s?', '', cover_lines[i])); i += 1
            block = [b for b in block if b.strip()]
            if block: callout(doc, block)
        else:
            i += 1

def portadilla(doc, titulo, subtitulo):
    doc.add_page_break()
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('ANEXO'); r.bold=True; r.font.name=FUENTE; r.font.size=Pt(12); r.font.color.rgb=CORAL
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo); r.bold=True; r.font.name=FUENTE; r.font.size=Pt(26); r.font.color.rgb=AZUL
    if subtitulo:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, subtitulo, size=12, color=GRIS)
        for r in p.runs: r.italic = True

# ===== Documento =====
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.0); s.left_margin = s.right_margin = Cm(2.2)
setup_footer(doc)

# 1) Portada + cuerpo de la Carpeta
carpeta_lines = open(CARPETA, encoding='utf-8').read().split('\n')
cover_lines, body_lines = split_cover(carpeta_lines)
render_cover(doc, cover_lines)

# 2) Índice (estático)
secciones = [l[3:].strip() for l in body_lines if l.startswith('## ')]
anexo_titulos = []
for fn in ANEXOS:
    first_line = open(f'{REPO}/{fn}', encoding='utf-8').readline().strip()
    anexo_titulos.append(re.sub(r'^#\s*', '', first_line).replace(' — ETC', ''))
doc.add_page_break()
h1(doc, 'Contenido')
for s in secciones:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    add_runs(p, re.sub(r'\*', '', s), size=11, color=TINTA)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
r = p.add_run('Anexos'); r.bold = True; r.font.name = FUENTE; r.font.size = Pt(11); r.font.color.rgb = CORAL
for k, t in enumerate(anexo_titulos, 1):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    add_runs(p, f'Anexo {k} · {t}', size=11, color=TINTA)

# 3) Cuerpo de la Carpeta (cada sección en su página)
doc.add_page_break()
parse_body(doc, body_lines, page_break_per_section=True)

# 4) Anexos (cada uno con portadilla; secciones fluyen)
for fn in ANEXOS:
    lines = open(f'{REPO}/{fn}', encoding='utf-8').read().split('\n')
    cov, body = split_cover(lines)
    titulo = re.sub(r'^#\s*', '', next((l for l in cov if l.startswith('# ')), 'Anexo')).replace(' — ETC', '')
    subt = next((l[4:] for l in cov if l.startswith('### ')), '')
    portadilla(doc, titulo, subt)
    # intro callouts del anexo
    intro = [re.sub(r'^\s*>\s?', '', l) for l in cov if l.lstrip().startswith('>')]
    intro = [x for x in intro if x.strip()]
    doc.add_paragraph()
    if intro: callout(doc, intro)
    parse_body(doc, body, page_break_per_section=False)

doc.save(OUT)
print(f"Wrote {OUT} — Carpeta + {len(ANEXOS)} anexos, con índice y numeración de página")
