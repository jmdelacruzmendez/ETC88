#!/usr/bin/env python3
import os as _os
_R = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
"""Maqueta de la CANTERA del ETC 88 — gente que NO va a participar (ni equipo ni
participante), organizada en: PADRINOS y candidatos a EQUIPOS AUXILIARES.

Fuente: data/cantera.json (el listado inicial de abril, leído del Drive) + data/equipo.json
(el equipo FINAL, para excluir a quien ya está dentro). Trazable: cada nombre viene del
Borrador; las etiquetas de experiencia vienen del cruce con ETC 78/79/85.

Salida: Cantera_ETC88.xlsx (Padrinos · Auxiliares · Participantes-pool · Pool-no-equipo).
"""
import json, unicodedata, difflib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = f'{_R}'
CAN = json.load(open(f'{REPO}/data/cantera.json'))
EQ = json.load(open(f'{REPO}/data/equipo.json'))

CREMA, MAR, TIERRA, SAFARI, AMBAR, CUERO = 'F7EBCC', '1B3A52', 'B25028', '4A5D2E', 'C57920', '6B4423'
thin = Side(style='thin', color='C9AE76'); border = Border(left=thin, right=thin, top=thin, bottom=thin)

def norm(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower()
    return [t for t in s.replace('/', ' ').split() if len(t) >= 2 and t not in ('de', 'la', 'del', 'los')]

# tokens de cada persona del equipo FINAL (todos: operativos, backups, vacante, ampliados)
roster = [norm(p['nombre'].replace(' (sin formulario)', '')) for p in EQ['equipo']]

def tok_match(a, b):
    if a == b: return True
    if len(a) >= 4 and len(b) >= 4 and (a[:4] == b[:4]): return True
    return difflib.SequenceMatcher(None, a, b).ratio() >= 0.86

def en_equipo(nombre):
    bt = norm(nombre)
    if not bt: return False
    for rt in roster:
        if not rt: continue
        # ¿coincide el primer nombre (o variante cercana)?
        first = tok_match(bt[0], rt[0])
        overlap = sum(1 for x in bt for y in rt if tok_match(x, y))
        if first and (overlap >= 2 or len(rt) == 1 or len(bt) == 1):
            return True
        if overlap >= 2 and (tok_match(bt[-1], rt[-1])):  # apellidos coinciden
            return True
    return False

def split_loc(entry):
    if '|' in entry:
        n, loc = entry.split('|', 1); return n.strip(), loc.strip()
    return entry.strip(), ''

EXP = CAN.get('experiencia_etc', {})
def exp_of(nombre):
    for k, v in EXP.items():
        if k == '_nota': continue
        if norm(k)[0] == norm(nombre)[0] and norm(k)[-1] == norm(nombre)[-1]:
            return v
    return ''

wb = openpyxl.Workbook(); wb.remove(wb.active)

def title(ws, txt, color=MAR, span=4):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(row=1, column=1, value=txt); c.font = Font(bold=True, size=13, color=CREMA)
    c.fill = PatternFill('solid', fgColor=color); c.alignment = Alignment('left', vertical='center')
    ws.row_dimensions[1].height = 26

def header(ws, r, cols, fill=TIERRA):
    for i, (n, w) in enumerate(cols, 1):
        c = ws.cell(row=r, column=i, value=n); c.font = Font(bold=True, size=10, color=CREMA)
        c.fill = PatternFill('solid', fgColor=fill); c.alignment = Alignment('center', vertical='center', wrap_text=True)
        c.border = border; ws.column_dimensions[get_column_letter(i)].width = w

def row(ws, r, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.font = Font(size=10); c.border = border
        c.alignment = Alignment('left', vertical='center', wrap_text=True)

# ---------------- 1. PADRINOS (no-misioneros + no-se-estiman, sin los que están en el equipo)
ws = wb.create_sheet('Padrinos')
title(ws, 'PADRINOS / APOYO · etecianos que NO sirven este año (donaciones · apoyo en Profondo #1 · oración)', CUERO, 4)
header(ws, 3, [('Nombre', 30), ('Ubicación', 14), ('Origen (Borrador 88)', 22), ('Nota', 30)], CUERO)
r = 4; n_pad = 0
for cat, etiqueta in [('no_misioneros', 'No-misionero/eteciano'), ('no_se_estiman', 'No se estima (reserva)')]:
    for entry in CAN[cat]:
        nombre, loc = split_loc(entry)
        if en_equipo(nombre):
            continue  # ya está en el equipo final (ej. Roberto Figueroa, Mary Carmen)
        row(ws, r, [nombre, loc, etiqueta, exp_of(nombre)]); r += 1; n_pad += 1

# ---------------- 2. AUXILIARES (mapeo curado de candidatos por equipo)
ws = wb.create_sheet('Auxiliares')
title(ws, 'CANDIDATOS A EQUIPOS AUXILIARES · de la cantera (a confirmar por Co-Dir)', SAFARI, 4)
header(ws, 3, [('Equipo auxiliar', 26), ('Candidatos (cantera, no-equipo)', 46), ('Por qué', 30), ('Estado', 12)], SAFARI)
AUX = [
    ('Recaudación y Donaciones',
     'Kedward Acevedo (Dir. Cocina ETC 85 · liderazgo) · Brissa Rodríguez · Carla Uribe · Lisset Rosario · Ramón Leonardo · Ángel Radesky · diáspora',
     'red/contactos para Profondo #1, donaciones y padrinazgo'),
    ('Actividad Profondo',
     'Mismo núcleo de Recaudación + apoyos externos de la cantera',
     'Profondo #1 es transversal; apoyos externos amplían alcance'),
    ('Guagua (Transporte)',
     'Etecianos con vehículo fuera del equipo (PC/SD): de "no se estiman" para tramos · Kedward Acevedo (PC, logística)',
     'rol externo, no pernocta en la casa'),
    ('Intersección (espiritual)',
     'Liderado por Padre Paul + Sor Angelina · sumar etecianos de oración de los no-misioneros',
     'acompañamiento/oración, no requiere estar en la casa'),
]
r = 4
for eq, cand, pq in AUX:
    row(ws, r, [eq, cand, pq, '[POR DEFINIR]']); r += 1
ws.cell(row=r+1, column=1, value='Nota: los responsables de cada equipo auxiliar los nombra la Co-Dirección (pendiente). '
        'Roberto, Guido, Franklin y Paloma ya están en el equipo y propusieron ideas de recaudación.').font = Font(italic=True, size=9)

# ---------------- 3. PARTICIPANTES (pool — SÍ participan; separado)
ws = wb.create_sheet('Participantes-pool')
title(ws, 'POOL DE PARTICIPANTES (del Borrador) · estos SÍ participan — no son cantera de apoyo', AMBAR, 3)
header(ws, 3, [('Nombre', 30), ('Origen', 20), ('Nota', 40)], AMBAR)
r = 4
for entry in CAN['participantes']:
    nombre, loc = split_loc(entry)
    row(ws, r, [nombre, 'Borrador 88 · participantes', 'cruzar con los invitados del formulario']); r += 1

# ---------------- 4. POOL no-equipo (guías/cocina/música del borrador que NO quedaron)
ws = wb.create_sheet('Pool-no-equipo')
title(ws, 'POOL: guías/cocina/música del Borrador que NO quedaron en el equipo final [aprox · verificar]', MAR, 4)
header(ws, 3, [('Nombre', 30), ('Área (borrador)', 16), ('Experiencia ETC', 26), ('Uso sugerido', 22)], MAR)
r = 4; n_pool = 0
for cat in ['guias', 'cocina', 'musica']:
    for entry in CAN[cat]:
        nombre, loc = split_loc(entry)
        if en_equipo(nombre):
            continue
        exp = exp_of(nombre)
        uso = 'Auxiliar/Profondo (con experiencia)' if exp else 'Padrino o participante'
        row(ws, r, [nombre + (f' ({loc})' if loc else ''), cat, exp, uso]); r += 1; n_pool += 1

wb.save(f'{REPO}/Cantera_ETC88.xlsx')
print(f"Wrote Cantera_ETC88.xlsx · padrinos {n_pad} · pool no-equipo {n_pool} · participantes {len(CAN['participantes'])}")
