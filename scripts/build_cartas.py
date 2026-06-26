#!/usr/bin/env python3
import os as _os
_R = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
"""Genera los documentos de recaudación del ETC 88 (.docx) + el tracker (CSV).

Por qué .docx: subir markdown/text a Drive escapa símbolos (\\-, \\[, \\#) y rompe
acentos/emoji; python-docx produce documentos que Google Docs importa limpios.

Formato CALCADO de las cartas reales de la comunidad (Diócesis de SPM · "Caminos de
Vida" · asociación de fieles): logo + membrete + cuerpo + firma de la Co-Dirección.

Salidas (raíz del repo):
  1. Carta_Induveca_ETC88.docx       — PRIMERA carta, personalizada (embutidos, fecha hoy).
  2. Carta_Modelo_Empresa_ETC88.docx — modelo empresa/negocio (general + por rubro).
  3. Carta_Modelo_Personal_ETC88.docx— modelo padrino/particular.
  4. Directorio_Donantes_ETC88.docx  — empresas/contactos/correos (SPM + nacionales).
  5. Donantes_Anteriores_ETC88.docx  — consolidado de quienes ya donaron (verificado).
  6. preparacion/TRACKER_CARTAS_DONACION_ETC88.csv — seguimiento de envíos.

Fuente: data/equipo.json (embebe data/estado.json) para fechas/número/Co-Dirección.
Contactos/donantes: preparacion/DIRECTORIO_DONANTES_ETC88.md (evidencia primaria Drive
ETC 78 + xlsx Misiones 2025 + investigación web). El logo se incrusta si existe
data/logo_etc.png (descargado del Drive de la comunidad: "Pez 1000x1000.png").
Teléfonos de la Co-Dirección dados por el director (18-jun-2026).
"""
import json, csv, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = f'{_R}'
LOGO = f'{REPO}/data/logo_etc.png'
FIRMA_SELLO = f'{REPO}/data/firma_sello.png'  # firma+sello reales (recorte del escaneo firmado)
D = json.load(open(f'{REPO}/data/equipo.json'))
EST = D['estado']

def _v(node):
    return node['valor'] if isinstance(node, dict) and 'valor' in node else node

R = EST.get('retiro', {})
NUM = _v(R.get('numero', 88))
FECHAS = _v(R.get('fechas', '4–6 de septiembre de 2026'))
CODIR = _v(R.get('co_direccion', ['Juan Manuel de la Cruz Méndez', 'Jean Carlo de la Cruz Méndez']))
if isinstance(CODIR, str):
    CODIR = [CODIR]
TELS = ['829-898-1416', '849-850-5178']  # Juan Manuel · Jean Carlo (Co-Dir, 18-jun-2026)
CUOTA = _v(EST.get('finanzas', {}).get('cuota_participante', 3000))
CUOTA_TXT = f"RD${CUOTA:,}" if isinstance(CUOTA, (int, float)) else f"RD${CUOTA}"
FECHA_HOY = '18 de junio de 2026'

MAR = RGBColor(0x1B, 0x3A, 0x52)
GRIS = RGBColor(0x80, 0x80, 0x80)
CEN = WD_ALIGN_PARAGRAPH.CENTER
DER = WD_ALIGN_PARAGRAPH.RIGHT

# ---------------------------------------------------------------- helpers ----
def new_doc():
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    return doc

def _p(doc, text='', italic=False, bold=False, align=None, size=None, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.italic, r.bold = italic, bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        rr = t.rows[0].cells[i].paragraphs[0].add_run(h)
        rr.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            if i < len(cells):
                cells[i].paragraphs[0].add_run('' if c is None else str(c))
    doc.add_paragraph()
    return t

def membrete(doc, fecha='[fecha]'):
    if os.path.exists(LOGO):
        p = doc.add_paragraph(); p.alignment = CEN
        p.add_run().add_picture(LOGO, width=Inches(1.5))
    else:
        _p(doc, '[ logo ETC (ΙΧΘΥΣ) · insertar ]', align=CEN, size=9, color=GRIS)
    _p(doc, 'DIÓCESIS DE SAN PEDRO DE MACORÍS', bold=True, align=CEN, size=13, color=MAR)
    _p(doc, 'Encuentro Total con Cristo (ETC) · Comunidad «Caminos de Vida»', italic=True, align=CEN, size=10)
    _p(doc, 'Asociación de fieles', align=CEN, size=9, color=GRIS)
    _p(doc, f'San Pedro de Macorís, Rep. Dom. — {fecha}', italic=True, align=DER)
    _p(doc)

def firma(doc):
    _p(doc)
    t = doc.add_table(rows=3, cols=2)  # tabla sin estilo = sin bordes
    n1 = CODIR[0] if CODIR else '[Co-Director]'
    n2 = CODIR[1] if len(CODIR) > 1 else ''
    t.rows[0].cells[0].paragraphs[0].add_run('____________________')
    t.rows[0].cells[1].paragraphs[0].add_run('____________________')
    r1 = t.rows[1].cells[0].paragraphs[0].add_run(n1); r1.bold = True
    r2 = t.rows[1].cells[1].paragraphs[0].add_run(n2); r2.bold = True
    t.rows[2].cells[0].paragraphs[0].add_run(TELS[0] if len(TELS) > 0 else '')
    t.rows[2].cells[1].paragraphs[0].add_run(TELS[1] if len(TELS) > 1 else '')
    _p(doc, f'Co-Dirección · Encuentro Total con Cristo (ETC) {NUM}')

def firma_imagen(doc):
    """Inserta la firma+sello reales (data/firma_sello.png, recorte del escaneo firmado).
    Solo con autorización de la Co-Dirección. Si falta el archivo, cae a la firma tipográfica."""
    _p(doc)
    if os.path.exists(FIRMA_SELLO):
        doc.add_paragraph().add_run().add_picture(FIRMA_SELLO, width=Inches(6.2))
    else:
        firma(doc)

SALUDO = ('Reciba un afectuoso saludo de nuestra parte y que la paz y el amor de Dios esté '
          'llenando cada espacio de su vida. Después de un cordial saludo en Cristo Jesús, '
          'hacemos de su conocimiento que el Encuentro Total con Cristo (ETC), asociación de '
          'fieles de la Iglesia Católica en la Diócesis de San Pedro de Macorís, convencidos de '
          'la necesidad de orientar y acompañar a los jóvenes de hoy por los caminos de Dios, '
          'trabaja en integrarlos con su familia humana y de fe para que realicen su vocación '
          'laical en el mundo y sean testimonio del amor de Cristo ante los demás.')

ACTIVIDAD = ('Realizamos diferentes actividades en grupo, entre las cuales la fundamental es '
             'nuestro retiro de un fin de semana llamado «Encuentro Total con Cristo (ETC)», en '
             'el que, a través de la reflexión y la experiencia de vida, se comparten los valores '
             f'cristianos. En esta ocasión realizaremos dicho encuentro del {FECHAS}, con cerca '
             'de 100 jóvenes.')

DESPEDIDA = ('Nos despedimos en espera de su respuesta y deseando que el Señor Jesús le colme '
             'de abundantes bendiciones.')

def destinatario(doc, lineas):
    for ln in lineas:
        _p(doc, ln)
    _p(doc, 'Sus manos,')
    _p(doc)

# ------------------------------------------------------------------ cartas ----
def carta_induveca(firmada=False):
    doc = new_doc()
    membrete(doc, FECHA_HOY)
    destinatario(doc, ['Señora', 'Daisy Medina', 'Induveca (Grupo SID)',
                       'Av. Máximo Gómez #182, Santo Domingo, R.D.'])
    _p(doc, SALUDO)
    _p(doc, ACTIVIDAD)
    _p(doc, 'Conocedores del compromiso social de Grupo SID con las familias dominicanas, '
            'apelamos a la generosidad de Induveca solicitando su colaboración, según su '
            'disponibilidad, con productos de su línea de embutidos y cárnicos (jamón, salami, '
            'salchichón, mortadela, salchichas), entre otros productos de su marca, para la '
            'alimentación de los jóvenes durante el fin de semana del retiro.')
    _p(doc, 'Si dentro del Grupo SID fuera posible, también nos ayudaría el aceite comestible '
            '(Crisol / Mazola) y la harina de MercaSID. Cualquier aporte, en producto o en '
            'efectivo, será de gran ayuda; por ser una asociación de fieles de la Diócesis de '
            'San Pedro de Macorís podemos emitir la constancia de su donación que requieran, y '
            'con gusto los incluimos en nuestra cadena de oración.')
    _p(doc, DESPEDIDA)
    (firma_imagen if firmada else firma)(doc)
    out = f'{REPO}/Carta_Induveca_FIRMADA_ETC88.docx' if firmada else f'{REPO}/Carta_Induveca_ETC88.docx'
    doc.save(out); print(f'Wrote {out}')

def carta_firmada(donante_lineas, out_name, fecha=FECHA_HOY, pedido=None):
    """Carta de donación CON firma+sello reales — cambian fecha, donante y (opcional) el pedido."""
    doc = new_doc()
    membrete(doc, fecha)
    destinatario(doc, donante_lineas)
    _p(doc, SALUDO); _p(doc, ACTIVIDAD)
    _p(doc, pedido or (
        'Para realizar este encuentro acudimos a la generosidad de instituciones y personas '
        'que colaboran con esta obra; por tal motivo le solicitamos que, según su posibilidad, '
        'nos haga una donación para este retiro, de forma que sirva de apoyo a nuestro '
        'presupuesto general. Por ser asociación de fieles de la Diócesis de San Pedro de '
        'Macorís podemos emitir la constancia de su donación; con gusto lo incluimos en '
        'nuestra cadena de oración.'))
    _p(doc, DESPEDIDA)
    firma_imagen(doc)
    doc.save(out_name); print(f'Wrote {out_name}')

def carta_personal_firmada(destino_lineas, out_name, fecha=FECHA_HOY):
    """Carta a padrino/particular CON firma+sello reales — solo cambia fecha y nombre."""
    doc = new_doc()
    membrete(doc, fecha)
    destinatario(doc, destino_lineas)
    _p(doc, SALUDO); _p(doc, ACTIVIDAD)
    _p(doc, 'Para que ningún joven se quede fuera por razones económicas, te invitamos a '
            'APADRINAR de una de estas formas:')
    for b in [f'Beca de un joven (su cuota): {CUOTA_TXT}',
              'Una Biblia o un pez con el nombre del padrino y del participante',
              'Un aporte libre, en el monto que decidas']:
        doc.add_paragraph(style='List Bullet').add_run(b)
    _p(doc, 'Puedes hacerlo por transferencia [cuenta — POR DEFINIR]. Con gusto te emitimos una '
            'constancia y te incluimos en nuestra cadena de oración y en la Misa de clausura.')
    _p(doc, DESPEDIDA)
    firma_imagen(doc)
    doc.save(out_name); print(f'Wrote {out_name}')

def carta_modelo_empresa():
    doc = new_doc()
    _p(doc, f'MODELO — Carta a Empresa / Negocio · ETC {NUM}', bold=True, align=CEN, size=16, color=MAR)
    _p(doc, 'Personaliza razón social, contacto y el pedido. Dos variantes abajo.', italic=True, align=CEN, size=10)
    doc.add_page_break()
    _p(doc, 'Variante 1 — Donación general', bold=True, color=MAR)
    membrete(doc)
    destinatario(doc, ['Señores', '[Razón social]',
                       "At'n.: [contacto / Depto. de Responsabilidad Social]", '[ciudad]'])
    _p(doc, SALUDO); _p(doc, ACTIVIDAD)
    _p(doc, 'Para realizar este encuentro acudimos a la generosidad de instituciones y personas '
            'que colaboran con esta obra; por tal motivo le solicitamos que, según su posibilidad, '
            'nos haga una donación para este retiro, de forma que sirva de apoyo a nuestro '
            'presupuesto general. Por ser asociación de fieles de la Diócesis de SPM podemos '
            'emitir la constancia de su donación; con gusto lo incluimos en nuestra cadena de oración.')
    _p(doc, DESPEDIDA); firma(doc)
    doc.add_page_break()
    _p(doc, 'Variante 2 — Por rubro (artículos del presupuesto)', bold=True, color=MAR)
    membrete(doc)
    destinatario(doc, ['Señores', '[Razón social]', "At'n.: [contacto]", '[ciudad]'])
    _p(doc, SALUDO); _p(doc, ACTIVIDAD)
    _p(doc, 'En tal sentido apelamos a su colaboración solicitando nos faciliten artículos que '
            'ustedes producen o comercializan y tenemos en nuestro presupuesto, tales como:')
    table(doc, ['Cantidad', 'Unidad', 'Artículo'], [['', '', ''] for _ in range(5)])
    _p(doc, DESPEDIDA); firma(doc)
    out = f'{REPO}/Carta_Modelo_Empresa_ETC88.docx'
    doc.save(out); print(f'Wrote {out}')

def carta_modelo_personal():
    doc = new_doc()
    _p(doc, f'MODELO — Carta a Padrino / Particular · ETC {NUM}', bold=True, align=CEN, size=16, color=MAR)
    doc.add_page_break()
    membrete(doc)
    destinatario(doc, ['Estimado(a) [nombre]', '[ciudad]'])
    _p(doc, SALUDO); _p(doc, ACTIVIDAD)
    _p(doc, 'Para que ningún joven se quede fuera por razones económicas, te invitamos a '
            'APADRINAR de una de estas formas:')
    for b in [f'Beca de un joven (su cuota): {CUOTA_TXT}',
              'Una Biblia o un pez con el nombre del padrino y del participante',
              'Un aporte libre, en el monto que decidas']:
        doc.add_paragraph(style='List Bullet').add_run(b)
    _p(doc, 'Puedes hacerlo por transferencia [cuenta — POR DEFINIR]. Con gusto te emitimos una '
            'constancia y te incluimos en nuestra cadena de oración y en la Misa de clausura.')
    _p(doc, DESPEDIDA); firma(doc)
    out = f'{REPO}/Carta_Modelo_Personal_ETC88.docx'
    doc.save(out); print(f'Wrote {out}')

# -------------------------------------------------------------- directorio ----
FUNDACIONES = [
    ['Fundación Grupo Puntacana', 'fgpc@puntacana.com · 809-959-2714', 'Comunidades del Este'],
    ['Fundación Central Romana (La Romana)', 'carta a RSE · centralromana.com.do', 'Salud/educación Este'],
    ['Voluntariado Banreservas', '809-960-2121 (entregar carta)', 'Útiles/aporte, nacional'],
    ['Grupo SID / MercaSID (RSE)', 'compras@mercasid.com.do · 809-565-2151', 'Dona producto'],
    ['Fundación Sur Futuro', 'surfuturo.org/contacto (formulario)', 'Desarrollo social'],
    ['UCE — Univ. Central del Este (SPM)', 'infopc@uce.edu.do · 809-933-1500', 'Institución local SPM'],
]
EMPRESAS = [
    ['Alimentos', 'Induveca (Grupo SID) — embutidos/pollo', 'servicioalcliente@induveca.com.do', 'Nacional'],
    ['Alimentos', 'Grupo SID / MercaSID (aceite Crisol)', 'compras@mercasid.com.do · 809-565-2151', 'Nacional'],
    ['Alimentos', 'Pollo Cibao (RSE)', 'r.sociales@pollocibao.com · 809-590-8520', 'Nacional'],
    ['Alimentos', 'Hipermercados Iberia (SPM)', '809-529-2799 · Av. Independencia 20', 'SPM'],
    ['Alimentos', 'César Iglesias S.A.', '[verificar] (Lic. Luis Velázquez)', 'SPM/Nacional'],
    ['Alimentos', 'Grupo Ramos (Sirena/Jumbo, RSE)', '809-472-4444 ext. 11215', 'Jumbo en SPM'],
    ['Biblias', 'Sociedad Bíblica Dominicana', '809-685-2025', 'Nacional (envío)'],
    ['Biblias', 'San Pablo / Paulinas / Cuesta / SDQLee', 'sanpablo.do · 809-685-7542', 'Nacional'],
    ['Impresión', 'Print Mate (SPM)', 'IG @printmatesrl · C/ Luís Amiama Tió', 'SPM'],
    ['Impresión', 'Editora y Papelería 23', '[verificar]', 'SPM'],
    ['Ferretería', 'Ferretería Serie 23 (SPM)', '809-529-7764', 'SPM'],
    ['Transporte', 'DominicanBus', '809-530-9742 · PC 829-946-3500 · info@dominicanbus.com', 'Nacional/Este'],
    ['Transporte', 'Metro Servicios · Caribe Tours', '809-530-2850 · 809-221-4422', 'Nacional'],
    ['Cofre/desechables', 'Multi Box (Sra. Gabriela Rodríguez)', 'SPM · [verificar]', 'SPM'],
]
PREMIOS = [
    ['Noche/estadía', 'Casa de Campo (La Romana)', '809-523-8171 · servicioalcliente@costasur.com.do'],
    ['Noche/estadía', 'Emotions by Hodelpa (Juan Dolio)', 'reservas.emotions@hodelpa.com · 809-683-3636'],
    ['Estadía', 'Bahía Príncipe — Grupos & Eventos', 'groups1@bahia-principe.com'],
    ['Electrodoméstico/gift card', 'Casa Cuesta · Plaza Lama · Bonos CCN', '809-537-5646 · 809-274-5262 · 809-544-5555'],
    ['Tour/excursión', 'Saona Dreams (Juan Dolio)', '809-556-1008 · resa@saonadreams.com'],
]

def directorio_doc():
    doc = new_doc()
    _p(doc, f'Directorio de Donantes — ETC {NUM}', bold=True, align=CEN, size=18, color=MAR)
    _p(doc, 'Empresas, contactos y correos (SPM + nacionales). Confirmar por WhatsApp/llamada '
            'antes de enviar; los marcados [verificar] aún no están confirmados.', italic=True, align=CEN, size=9)
    _p(doc)
    _p(doc, 'Fundaciones / Responsabilidad Social', bold=True, color=MAR)
    table(doc, ['Fundación', 'Contacto', 'Enfoque'], FUNDACIONES)
    _p(doc, 'Empresas por rubro', bold=True, color=MAR)
    table(doc, ['Rubro', 'Empresa', 'Contacto / correo', 'Zona'], EMPRESAS)
    _p(doc, 'Premios de rifa (donados)', bold=True, color=MAR)
    table(doc, ['Premio', 'Donante candidato', 'Contacto'], PREMIOS)
    out = f'{REPO}/Directorio_Donantes_ETC88.docx'
    doc.save(out); print(f'Wrote {out}')

# --------------------------------------------------------- donantes previos ---
DON_EMPRESAS_2025 = [
    ['SEDESTE Group', 'Efectivo', 'RD$10,000'],
    ['Agys Ferreservis (Steven García)', 'Efectivo', 'RD$10,000'],
    ['Paulino Ingeniería', 'Efectivo', 'RD$5,000'],
    ['Zaglul', 'Cheque', 'RD$3,000'],
]
DON_78 = [
    ['Jessica de León', 'Efectivo', 'RD$2,000'],
    ['Melissa Lugo', 'Efectivo', 'RD$3,000'],
    ['Carolina Castillo', 'Efectivo', 'RD$1,000'],
    ['"Manolito"', 'Efectivo', 'RD$2,000'],
    ['Samuel Humphry', 'Efectivo', 'RD$60'],
    ['Etecianos (equipo)', 'Especie', 'Ofrenda sábado'],
    ['Iberia', 'Especie (probable)', '5 víveres de cocina'],
]
ESPECIE_2025 = ('Luis Manuel (agua), P. Paul (arroz/leche), P. Ligonde (botellones), Nelson '
                '(azúcar), Victoria (café), María Astacio (canela/chocolate/cloro), Belkis '
                '(harina), Ricaira (mayonesa), Rissel (huevos), Wirna (leche), Yulainy (batata), '
                'Maritza (salami Induveca), Kamila (caldo), Raiza (desechables), Juan Solano '
                '(servilletas), Camila Acta (detergente), Mamá Mariangel (lápices), Robiladi '
                '(dulces), Madallyn (cartulinas).')

def donantes_doc():
    doc = new_doc()
    _p(doc, f'Donantes Anteriores (consolidado) — hacia ETC {NUM}', bold=True, align=CEN, size=18, color=MAR)
    _p(doc, 'Solo lo verificado en documentos primarios (xlsx Misiones 2025 + Drive ETC 78). '
            'ETC 85 no dejó registro de donantes (solo solicitudes).', italic=True, align=CEN, size=9)
    _p(doc)
    _p(doc, 'Misiones Semana Santa 2025 — empresas/instituciones', bold=True, color=MAR)
    table(doc, ['Donante', 'Tipo', 'Monto'], DON_EMPRESAS_2025)
    _p(doc, 'Individuos (efectivo): ~50 aportantes, total RD$199,537 (+ Profondo RD$64,375 + '
            'cupos RD$70,000 = RD$333,912). Lista nominal completa en el xlsx de Misiones 2025.', size=10)
    _p(doc, 'Misiones 2025 — donantes en especie (cocina/limpieza)', bold=True, color=MAR)
    _p(doc, ESPECIE_2025, size=10)
    _p(doc)
    _p(doc, 'ETC 78 — verificado en presupuesto/ingresos', bold=True, color=MAR)
    table(doc, ['Donante', 'Tipo', 'Detalle'], DON_78)
    _p(doc, 'Donantes recurrentes (78 + 2025): Melissa Lugo, Giselle Núñez, Zaglul, María '
            'Astacio, Belkis — los más cálidos para el ETC 88.', italic=True, size=10)
    out = f'{REPO}/Donantes_Anteriores_ETC88.docx'
    doc.save(out); print(f'Wrote {out}')

# ---------------------------------------------------------------- tracker ----
TARGETS = [
    ['1', 'Induveca (Grupo SID)', 'Alimentos (embutidos)', 'Empresa-rubro', 'servicioalcliente@induveca.com.do', 'web·verificar · PRIMERA carta enviada', 'salami, jamón, salchichón, mortadela'],
    ['1', 'Grupo SID / MercaSID', 'Alimentos (aceite/harina)', 'Empresa-rubro', 'compras@mercasid.com.do · 809-565-2151', 'web·verificar (dona producto)', 'aceite Crisol, harina'],
    ['1', 'César Iglesias S.A. (Lic. Luis Velázquez)', 'Alimentos/limpieza', 'Empresa-rubro', 'SPM · [verificar]', 'SOLICITADO 2023 y 2025 (repetida)', 'aceite, pasta, cloro, papel'],
    ['1', 'Almacenes Iberia (Sr. Alberto Rivera)', 'Alimentos', 'Empresa-rubro', '809-529-2799 (SPM)', 'SOLICITADO 78 · especie probable [verificar]', 'lácteos, pan, víveres'],
    ['1', 'EGE Haina (Resp. Social)', 'Corporativo/RSE', 'Empresa', '[web·verificar]', 'SOLICITADO 78 (vía RSE)', 'aporte económico'],
    ['1', 'Almacenes Zaglul (Dr. Ramón Zaglul)', 'Alimentos/varios', 'Empresa-rubro', '[verificar]', 'SOLICITADO 78 · DONÓ RD$3,000 (2025)', 'víveres / aporte'],
    ['2', 'SEDESTE Group', 'Corporativo', 'Empresa', '[verificar]', 'VERIFICADO 2025 (RD$10,000)', 'aporte económico'],
    ['2', 'Agys Ferreservis (Steven García)', 'Ferretería', 'Empresa', '[verificar]', 'VERIFICADO 2025 (RD$10,000)', 'aporte / ferretería'],
    ['2', 'Paulino Ingeniería', 'Corporativo', 'Empresa', '[verificar]', 'VERIFICADO 2025 (RD$5,000)', 'aporte económico'],
    ['3', 'Pollo Cibao', 'Alimentos (pollo)', 'Empresa-rubro', 'r.sociales@pollocibao.com', 'web·verificar (RSE)', 'pollo'],
    ['4', 'Casa de Campo (La Romana)', 'Premio de rifa', 'Empresa', '809-523-8171', 'web·verificar', 'noche/estadía'],
    ['4', 'Emotions by Hodelpa (Juan Dolio)', 'Premio de rifa', 'Empresa', 'reservas.emotions@hodelpa.com', 'web·verificar', 'noche/estadía'],
    ['5', 'Fundación Grupo Puntacana', 'Fundación', 'Empresa', 'fgpc@puntacana.com', 'web·verificar', 'proyecto / estadía'],
    ['5', 'Voluntariado Banreservas', 'Fundación', 'Empresa', '809-960-2121', 'web·verificar', 'útiles / aporte'],
    ['6', 'Sociedad Bíblica Dominicana', 'Biblias', 'Empresa-rubro', '809-685-2025', 'web·verificar', '50 biblias'],
    ['6', 'Print Mate (SPM)', 'Impresión', 'Empresa-rubro', 'IG @printmatesrl', 'SOLICITADO/usado 78', 'talonarios, cancionero'],
]

def write_tracker():
    path = f'{REPO}/preparacion/TRACKER_CARTAS_DONACION_ETC88.csv'
    cols = ['Prioridad', 'Destinatario', 'Rubro', 'Plantilla', 'Contacto', 'Procedencia',
            'Qué pedir', 'Fecha envío', 'Estado', 'Monto comprometido', 'Constancia',
            'Dueño del hilo', 'Próxima acción']
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(cols)
        for t in TARGETS:
            w.writerow(t + ['', '', '', 'No', '', ''])
    print(f'Wrote {path}')

def main():
    carta_induveca(firmada=False)
    carta_induveca(firmada=True)
    carta_modelo_empresa()
    carta_modelo_personal()
    carta_personal_firmada(['Estimado(a) [nombre]', '[ciudad]'], f'{REPO}/Carta_Modelo_Personal_FIRMADA_ETC88.docx')
    directorio_doc()
    donantes_doc()
    write_tracker()

if __name__ == '__main__':
    import sys
    if len(sys.argv) >= 3 and sys.argv[1] == 'firmada':
        # On-demand: python scripts/build_cartas.py firmada "Razón social" ["At'n. nombre"]
        donante = sys.argv[2]
        atn = sys.argv[3] if len(sys.argv) > 3 else None
        lineas = ['Señores', donante] + ([f"At'n.: {atn}"] if atn else [])
        slug = ''.join(c if c.isalnum() else '_' for c in donante).strip('_')[:40]
        carta_firmada(lineas, f'{REPO}/Carta_Firmada_{slug}.docx')
    elif len(sys.argv) >= 3 and sys.argv[1] == 'firmada-personal':
        # On-demand particular: python scripts/build_cartas.py firmada-personal "Nombre"
        nombre = sys.argv[2]
        slug = ''.join(c if c.isalnum() else '_' for c in nombre).strip('_')[:40]
        carta_personal_firmada([f'Estimado(a) {nombre}', '[ciudad]'], f'{REPO}/Carta_Firmada_Personal_{slug}.docx')
    else:
        main()
