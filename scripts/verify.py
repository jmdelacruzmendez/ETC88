#!/usr/bin/env python3
import os as _os
_R = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
"""COMPUERTA DE VERIFICACIÓN — se corre SIEMPRE antes de entregar algo del ETC 88.

Bloquea la entrega si algo no cuadra: conteos, nombres inventados, cifras sin
fuente, propuestas sin etiquetar, cadenas prohibidas, HTML roto o salidas viejas.
Imprime PASS/FAIL por chequeo y un resumen final; sale con código != 0 si hay FAIL.

Uso:  python scripts/verify.py
"""
import json, os, re, subprocess, sys
from docx import Document
import openpyxl

REPO = f'{_R}'
results = []  # (nombre, ok, detalle)

def check(nombre, ok, detalle=''):
    results.append((nombre, bool(ok), detalle))

# --------------------------------------------------------------- utilidades --
def docx_text(path):
    d = Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(out)

def xlsx_text(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(' '.join(str(c) for c in row if c is not None))
    return '\n'.join(out)

def read(path):
    return open(path, encoding='utf-8').read()

SHAREABLE_DOCX = ['Documento_Asesores_ETC88.docx', 'Carpeta_F1_ETC88.docx']
SHAREABLE_XLSX = ['Finanzas_ETC88.xlsx', 'Equipo_ETC88.xlsx']
SHAREABLE_MD   = ['GUIA_ETC88.md', 'GUIA_DE_GUIAS_88.md', 'ANEXO_COCINA_88.md']
SHAREABLE_HTML = ['index.html']

# ---------------------------------------------------- 1. JSON válido ---------
try:
    EQ = json.load(open(f'{REPO}/data/equipo.json'))
    EST = json.load(open(f'{REPO}/data/estado.json'))
    check('1. JSON válido (equipo.json + estado.json)', True)
except Exception as e:
    check('1. JSON válido (equipo.json + estado.json)', False, str(e))
    # sin JSON no se puede seguir
    EQ, EST = {'equipo': []}, {}

equipo = EQ.get('equipo', [])

# ---------------------------------------------------- 2. Conteos -------------
def count_area(area):
    return sum(1 for p in equipo if p['area'] == area
               and p.get('operativo') and not p.get('vacante') and not p.get('backup'))
try:
    ce = EST['conteos_esperados']
    errs = []
    if len(equipo) != ce['total_equipo']:
        errs.append(f"total {len(equipo)}≠{ce['total_equipo']}")
    nop = sum(1 for p in equipo if p.get('operativo'))
    if nop != ce['operativos']:
        errs.append(f"operativos {nop}≠{ce['operativos']}")
    for area, exp in ce['operativos_titulares_por_area'].items():
        got = count_area(area)
        if got != exp:
            errs.append(f"{area} {got}≠{exp}")
    nvac = sum(1 for p in equipo if p.get('vacante'))
    nbk = sum(1 for p in equipo if p.get('backup'))
    if nvac != ce['vacantes']:
        errs.append(f"vacantes {nvac}≠{ce['vacantes']}")
    if nbk != ce['backups']:
        errs.append(f"backups {nbk}≠{ce['backups']}")
    check('2. Conteos del roster = esperados', not errs, '; '.join(errs))
except Exception as e:
    check('2. Conteos del roster = esperados', False, str(e))

# ---------------------------------------------------- 3. Nombres -------------
# Conjunto de nombres reales del roster (completos) + tokens conocidos válidos.
roster_full = {p['nombre'].replace(' (sin formulario)', '').strip() for p in equipo}
guias_reales = sorted(p['nombre'] for p in equipo
                      if p['area'] == 'guias' and p.get('operativo')
                      and not p.get('backup') and not p.get('vacante'))
try:
    errs = []
    for f in SHAREABLE_DOCX:
        txt = docx_text(f'{REPO}/{f}')
        # 3a. token prohibido conocido
        if 'Dionis' in txt:
            errs.append(f"{f}: aparece 'Dionis' (no está en el equipo)")
        # 3b. Wirna NO puede salir como guía
        for m in re.finditer('Wirna', txt):
            ventana = txt[max(0, m.start() - 60):m.start() + 60]
            if re.search(r'gu[ií]a', ventana, re.I):
                errs.append(f"{f}: 'Wirna' aparece en contexto de guía (es Cocina)")
                break
    # 3c. La lista de guías del Documento de Asesores = los 14 reales
    txt_as = docx_text(f'{REPO}/Documento_Asesores_ETC88.docx')
    m = re.search(r'Guías \((\d+)\)', txt_as)
    n_guias = int(m.group(1)) if m else -1
    if n_guias != len(guias_reales):
        errs.append(f"Doc Asesores dice Guías ({n_guias}); el roster tiene {len(guias_reales)}")
    check('3. Sin nombres inventados (Dionis/Wirna-guía) + guías = roster', not errs, '; '.join(errs))
except Exception as e:
    check('3. Sin nombres inventados', False, str(e))

# ---------------------------------------------------- 4. Cifras-decisión -----
# Las cifras confirmadas deben aparecer; las viejas inventadas NO.
try:
    ef = EST['finanzas']
    txt_as = docx_text(f'{REPO}/Documento_Asesores_ETC88.docx')
    errs = []
    for label, val in [('casa sin exención', ef['casa_por_persona_sin_exencion']['valor']),
                       ('casa con exención', ef['casa_por_persona_con_exencion']['valor']),
                       ('deuda', ef['deuda_inicial']['valor']),
                       ('cuota participante', ef['cuota_participante']['valor'])]:
        if f"{val:,}" not in txt_as:
            errs.append(f"falta cifra confirmada {label}={val:,} en Doc Asesores")
    # Cifras inventadas viejas prohibidas en cualquier salida compartible
    BAD = ['1,250', '× $5,000', '5,000 mensual', '$5,000 mensual', '260,000', '260000']
    pool = txt_as + docx_text(f'{REPO}/Carpeta_F1_ETC88.docx') \
        + xlsx_text(f'{REPO}/Finanzas_ETC88.xlsx')
    for b in BAD:
        if b in pool:
            errs.append(f"aparece cifra inventada vieja: '{b}'")
    check('4. Cifras-decisión cuadran con estado.json (sin inventos viejos)', not errs, '; '.join(errs))
except Exception as e:
    check('4. Cifras-decisión cuadran con estado.json', False, str(e))

# ---------------------------------------------------- 5. Propuestas marcadas -
try:
    txt_as = docx_text(f'{REPO}/Documento_Asesores_ETC88.docx')
    fin_txt = xlsx_text(f'{REPO}/Finanzas_ETC88.xlsx')
    errs = []
    # 14-jun: cuota equipo ahora CONFIRMADA (2,000 total, era PROPUESTA 1,500–2,000).
    # Reglas vivas: lo no-confirmado en estado.json debe ir con [PROPUESTA] en docs.
    import json as _json
    _est = _json.load(open(f'{REPO}/data/estado.json', encoding='utf-8'))
    _ce_estado = _est['finanzas']['cuota_equipo'].get('estado', '')
    for needle in ['Cuota del equipo']:
        line = next((l for l in txt_as.split('\n') if needle in l), '')
        if _ce_estado.startswith('propuesta'):
            if '[PROPUESTA]' not in line:
                errs.append(f"'{needle}' (propuesta en estado.json) sin [PROPUESTA] en Doc Asesores")
        # Si está confirmada, no exigir [PROPUESTA] (sería incongruente).
    # En Finanzas, marcar PROPUESTA solo si HAY algo en estado.json todavía como propuesta
    has_proposals = any(
        isinstance(v, dict) and v.get('estado', '').startswith('propuesta')
        for v in _est['finanzas'].values()
    )
    if has_proposals and 'PROPUESTA' not in fin_txt:
        errs.append("Finanzas no marca ninguna PROPUESTA (estado.json tiene propuestas vivas)")
    check('5. Propuestas etiquetadas [PROPUESTA]', not errs, '; '.join(errs))
except Exception as e:
    check('5. Propuestas etiquetadas', False, str(e))

# ---------------------------------------------------- 6. Cadenas prohibidas --
FORBIDDEN = {
    'veterano': re.compile(r'veterano', re.I),
    'laico operativo': re.compile(r'laico operativo', re.I),
    'Constelación': re.compile(r'constelaci', re.I),
    'emoji roto (ð)': re.compile('ð'),
    'markdown escapado': re.compile(r'\\[-\[\]#=.]'),
}
try:
    errs = []
    pools = {}
    for f in SHAREABLE_HTML + SHAREABLE_MD:
        pools[f] = read(f'{REPO}/{f}')
    for f in SHAREABLE_DOCX:
        pools[f] = docx_text(f'{REPO}/{f}')
    for f in SHAREABLE_XLSX:
        pools[f] = xlsx_text(f'{REPO}/{f}')
    for fname, content in pools.items():
        for label, rx in FORBIDDEN.items():
            if rx.search(content):
                errs.append(f"{fname}: '{label}'")
    check('6. Cadenas prohibidas ausentes (veterano/laico op./Constelación/ð/\\md)',
          not errs, '; '.join(errs))
except Exception as e:
    check('6. Cadenas prohibidas ausentes', False, str(e))

# ---------------------------------------------------- 7. HTML + JS -----------
try:
    html = read(f'{REPO}/index.html')
    errs = []
    # 7a. balance de tags críticos
    for tag in ['div', 'script', 'template']:
        o = len(re.findall(fr'<{tag}[\s>]', html))
        c = len(re.findall(fr'</{tag}>', html))
        if o != c:
            errs.append(f"<{tag}> {o} ≠ </{tag}> {c}")
    # 7b. node --check del JS embebido (solo scripts JS, no application/json)
    blocks = re.findall(r'<script\b([^>]*)>(.*?)</script>', html, re.S)
    js = []
    for attrs, body in blocks:
        if 'src=' in attrs:
            continue
        if re.search(r'type\s*=\s*["\']application/(ld\+)?json["\']', attrs):
            continue
        js.append(body)
    if js:
        tmp = '/tmp/_verify_inline.js'
        open(tmp, 'w').write('\n;\n'.join(js))
        r = subprocess.run(['node', '--check', tmp], capture_output=True, text=True)
        if r.returncode != 0:
            errs.append('node --check falló: ' + (r.stderr.strip().split('\n')[0] if r.stderr else '?'))
    check('7. HTML balanceado + JS embebido válido (node --check)', not errs, '; '.join(errs))
except Exception as e:
    check('7. HTML + JS', False, str(e))

# ---------------------------------------------------- 8. Frescura ------------
def mtime(p):
    return os.path.getmtime(f'{REPO}/{p}') if os.path.exists(f'{REPO}/{p}') else 0
try:
    errs = []
    # equipo.json debe ser >= estado.json (se regeneró tras cambiar estado)
    if mtime('data/equipo.json') < mtime('data/estado.json'):
        errs.append('data/equipo.json más viejo que estado.json (corre build_data.py)')
    # cada salida >= su generador
    pairs = [
        ('index.html', 'scripts/build_html.py'),
        ('GUIA_ETC88.md', 'scripts/build_guia.py'),
        ('Finanzas_ETC88.xlsx', 'scripts/build_finanzas.py'),
        ('Equipo_ETC88.xlsx', 'scripts/build_excel.py'),
        ('Documento_Asesores_ETC88.docx', 'scripts/build_docx.py'),
        ('Carpeta_F1_ETC88.docx', 'scripts/build_docx.py'),
    ]
    for out, src in pairs:
        if mtime(out) < mtime(src):
            errs.append(f"{out} más viejo que {src} (regenera)")
        # y >= estado.json (todas las salidas dependen de los datos)
        if mtime(out) < mtime('data/estado.json'):
            errs.append(f"{out} más viejo que data/estado.json (regenera)")
    check('8. Frescura: salidas más nuevas que fuentes', not errs, '; '.join(errs))
except Exception as e:
    check('8. Frescura', False, str(e))

# ------- 9. Sin asignación-persona-rol sin respaldo de estado.json ----------
# Detecta patrones "[nombre propio] + palabra de rol" en docs de trabajo
# que NO están respaldados como confirmado en estado.json.
# Lista blanca de asignaciones CONFIRMADAS (fuente: estado.json equipos/asesores).
CONFIRMED_ROLE_PAIRS = {
    # (fragmento_nombre, keyword) — lo que SÍ puede aparecer
    ('Juan Manuel', 'director'), ('Jean Carlo', 'director'),
    ('Juan Manuel', 'Co-Dir'), ('Jean Carlo', 'Co-Dir'),
    ('Laura', 'asesora'), ('Tomás', 'asesor'), ('Frank', 'asesor'),
    ('Frank', 'Banderín'), ('Frank', 'banderín'),
    ('Priscilla', 'coord'), ('Camila', 'coord'),
    ('Paloma', 'coord'), ('Jhonnito', 'coord'),
    ('José Ángel', 'Música'), ('José Tusen', 'Música'),
    ('Paul', 'asesor'), ('Paul', 'transversal'),
    ('Sor Angelina', 'transversal'), ('Sor', 'transversal'),
    ('Leticia', 'asesor'), ('Marleny', 'asesor'), ('Sandrita', 'asesor'),
    ('Tomás', 'timekeeper'), ('Laura', 'Acta'),
    ('Samuel Montilla', 'administración'), ('Samuel Montilla', 'casa'),
}
# Nombres que NO deben aparecer con palabras de rol sin [POR DEFINIR] / [PROPUESTA]
ROLE_WORDS = re.compile(
    r'\b(l[ií]der|lidera|responsable|encargado|encargada|coordina|dirige|gestiona)\b', re.I
)
FLAG_NAMES = ['Roberto Figueroa', 'Guido Maldonado', 'Guido', 'Randolph',
              'Franklin', 'Kedward']
try:
    errs = []
    prep_files = [os.path.join(REPO, 'preparacion', f)
                  for f in os.listdir(os.path.join(REPO, 'preparacion'))
                  if f.endswith('.md') and f != 'TRAZABILIDAD.md']
    for fpath in prep_files:
        try:
            content = open(fpath, encoding='utf-8').read()
        except Exception:
            continue
        for name in FLAG_NAMES:
            for m in re.finditer(re.escape(name), content):
                # ventana de ±120 chars alrededor del nombre
                start = max(0, m.start() - 120)
                end = min(len(content), m.end() + 120)
                window = content[start:end]
                if ROLE_WORDS.search(window):
                    # ¿Está etiquetado como [POR DEFINIR] o [PROPUESTA]?
                    if '[POR DEFINIR]' not in window and '[PROPUESTA]' not in window:
                        # ¿Es un par confirmado?
                        confirmed = any(
                            n in window and k in window
                            for n, k in CONFIRMED_ROLE_PAIRS
                            if n in name or name in n
                        )
                        if not confirmed:
                            lineno = content[:m.start()].count('\n') + 1
                            fname = os.path.basename(fpath)
                            errs.append(
                                f"{fname}:{lineno}: '{name}' + rol sin [POR DEFINIR]/[PROPUESTA]"
                            )
                            break
    check('9. Sin asignación-persona-rol sin respaldo (formulario ≠ rol)',
          not errs, '; '.join(errs))
except Exception as e:
    check('9. Sin asignación-persona-rol sin respaldo', False, str(e))

# ---------------------------------------------------- resumen ----------------
print()
print('=' * 66)
print('  VERIFICACIÓN ETC 88')
print('=' * 66)
allok = True
for nombre, ok, detalle in results:
    print(f"  [{'PASS' if ok else 'FAIL'}] {nombre}")
    if not ok and detalle:
        print(f"         → {detalle}")
    allok = allok and ok
print('=' * 66)
print(f"  RESULTADO: {'TODO PASS ✅ — listo para entregar' if allok else 'HAY FALLOS ❌ — NO entregar'}")
print('=' * 66)
sys.exit(0 if allok else 1)
